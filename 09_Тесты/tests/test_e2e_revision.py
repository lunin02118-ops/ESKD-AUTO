# -*- coding: utf-8 -*-
"""E2E, группа V — кнопка «Новая ревизия» (ТЗ-02 Т-48…Т-52): штамп, журнал изменений и выгрузка с `_ИзмN`.

Изделие «выдаётся» отчётом `_Выдано_<дата>.txt` — так же, как его пишет кнопка «Выдать в производство»:
до выдачи чертёж правят свободно и ревизия не нужна.
"""
import hashlib
import os
import shutil
import stat
import unittest
from pathlib import Path

import openpyxl

from eskd_e2e import com, oracles, paths
from eskd_e2e.testing import SwTestCase

PRODUCT = "И01_ПРТИ.468211.100"
ASM = "ПРТИ.468211.100 СБ Кондуктор сварочный.sldasm"
PART = "ПРТИ.468211.101 Пластина опорная.sldprt"
DRAWING = "ПРТИ.468211.101 Пластина опорная.slddrw"


class Revision(SwTestCase):

    def _product(self, issued=True):
        """Изделие заказа с чертежами; issued — изделие уже выдано в производство."""
        short = self._case_name().split("_")[1]
        subdir = f"{short}/_Заявки/2026-001/02_Металл/{PRODUCT}/01_3D"
        models = self.s.run_dir / subdir
        # Каталог теста — целиком: журнал и отчёты прежнего прогона в папке изделия иначе остаются.
        if (self.s.run_dir / short).exists():
            shutil.rmtree(self.s.run_dir / short, ignore_errors=True)
        for src in sorted(Path(paths.FIXTURES_A).iterdir()):
            if src.suffix.lower() in (".sldprt", ".sldasm", ".slddrw"):
                self.s.workspace_copy(src, subdir=subdir)
        product = models.parent
        if issued:
            (product / "_Выдано_2026-09-10.txt").write_text(
                "Выдано в производство\n\nab12cd  " + PART + "\n", encoding="utf-8")
        return product, models

    def _revision(self, what="толщина 2 → 3 мм", code="4", backlog="доработать"):
        com.call(self.s.eskd(), "NewRevisionSilent", what, code, backlog)
        return str(com.call(self.s.eskd(), "RevisionStatus"))

    def test_V01_raises_revision_and_writes_journal(self):
        """V01: на чертеже выданной детали ревизия 0 → 1, строка в Изменения.xlsx, свойство Revision записано."""
        product, models = self._product()
        doc = self.s.open(models / DRAWING)
        self.s.activate(doc)
        status = self._revision()
        self.assertTrue(status.startswith("ok|"), status)
        _, revision, line, exported = status.split("|")
        self.assertEqual("1", revision, f"первая ревизия: {status}")
        self.assertEqual("1", line, f"первая строка журнала: {status}")

        journal = product / "Изменения.xlsx"
        self.assertTrue(journal.is_file(), "журнал изменений заведён")
        sheet = openpyxl.load_workbook(journal)["Изменения"]
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual("№", rows[0][0], f"шапка журнала: {rows[0]}")
        self.assertEqual("1", str(rows[1][1]), f"ревизия в журнале: {rows[1]}")
        self.assertEqual(DRAWING, rows[1][4], f"документ в журнале: {rows[1]}")
        self.assertEqual("толщина 2 → 3 мм", rows[1][5], f"что изменено: {rows[1]}")
        self.assertEqual("устранение ошибок", rows[1][6], f"причина по коду 4: {rows[1]}")
        self.assertEqual("доработать", rows[1][8], f"задел: {rows[1]}")

        self.s.close_all()
        self.assertEqual("1", oracles.value(self.persisted(models / DRAWING), "Revision"),
                         "ревизия записана в свойство чертежа")
        self.assertEqual([], self.addin_errors(), "ошибки в журнале надстройки")

    def test_V02_exports_with_revision_suffix_and_archives_previous(self):
        """V02: после ревизии PDF выгружается с суффиксом _Изм1, а прежний файл уходит в _Аннулировано (Т-30, Т-52)."""
        # Сначала черновик: выгрузка даёт PDF без суффикса — то, что уйдёт цеху…
        product, models = self._product(issued=False)
        doc = self.s.open(models / ASM)
        self.s.activate(doc)
        com.call(self.s.eskd(), "ExportProductSilent")
        self.assertTrue(str(com.call(self.s.eskd(), "ExportStatus")).startswith("ok|"), "черновая выгрузка")
        before = sorted(p.name for p in (product / "02_PDF").glob("ПРТИ.468211.101*.pdf"))
        self.assertTrue(before, "PDF детали до ревизии")
        self.s.close_all()
        # …и только потом выдача: выданное правят через ревизию.
        (product / "_Выдано_2026-09-10.txt").write_text(
            "Выдано в производство\n\nab12cd  " + PART + "\n", encoding="utf-8")

        drawing = self.s.open(models / DRAWING)
        self.s.activate(drawing)
        self.assertTrue(self._revision().startswith("ok|"))
        after = sorted(p.name for p in (product / "02_PDF").glob("ПРТИ.468211.101*.pdf"))
        self.assertTrue(any(name.endswith("_Изм1.pdf") for name in after), f"PDF новой ревизии: {after}")
        archived = sorted(p.name for p in (product / "02_PDF" / "_Аннулировано").glob("*.pdf"))
        self.assertTrue(archived, "прежний PDF перенесён в _Аннулировано")

    def test_V05_bch_part_revision_reexports_and_archives_only_its_files(self):
        """V05 (аудит 23.09.2026, NAME-1, REV-2): у детали БЧ ревизия — её свойство «Ревизия», а выгрузка читала только
        «Revision»: после новой ревизии выданная БЧ-деталь пропускалась («оформите новую ревизию»), а её прежняя развёртка
        уже лежала в «_Аннулировано» — у цеха не оставалось ничего. Теперь развёртка выгружается с «_Изм1», прежняя
        уходит в архив только после этого, а файл другого документа с похожим именем («… усиленная») остаётся на месте."""
        from eskd_e2e import build
        short = self._case_name().split("_")[1]
        if (self.s.run_dir / short).exists():
            shutil.rmtree(self.s.run_dir / short, ignore_errors=True)
        models = self.s.run_dir / f"{short}/_Заявки/2026-001/02_Металл/И01_ПРТИ.468211.190/01_3D"
        models.mkdir(parents=True)
        product = models.parent
        part = models / "ПРТИ.468211.191 Косынка.sldprt"
        doc = build.sheet_metal_plate(self.s, 200, 100, 3, "Лист 3,0 ГОСТ 19903-2015 / Ст3сп ГОСТ 16523-97")
        self.s.save_as(doc, part)
        self.s.activate(doc)
        self.assertEqual(1, int(com.call(self.s.eskd(), "ToggleDrawinglessSilent")), "деталь стала БЧ")
        self.s.save(doc)

        # Черновая выгрузка — развёртка без суффикса; рядом файл другого документа с тем же началом имени.
        com.call(self.s.eskd(), "ExportProductSilent")
        self.assertTrue(str(com.call(self.s.eskd(), "ExportStatus")).startswith("ok|"), "черновая выгрузка")
        drafts = [p for p in (product / "03_ЧПУ").rglob("*.dxf") if "_Аннулировано" not in p.parts]
        self.assertEqual(1, len(drafts), f"развёртка черновика: {drafts}")
        draft = drafts[0]
        self.assertNotIn("_Изм", draft.name, draft.name)
        decoy = draft.parent / "ПРТИ.468211.191 Косынка усиленная_S3мм_100х50.dxf"
        decoy.write_bytes(draft.read_bytes())
        (product / "_Выдано_2026-09-10.txt").write_text(
            "Выдано в производство\n\nab12cd  " + part.name + "\n", encoding="utf-8")

        self.s.activate(doc)
        status = self._revision()
        self.assertTrue(status.startswith("ok|"), status)
        self.assertEqual("1", status.split("|")[1], status)
        self.assertNotIn("остались файлы прежней ревизии", status, status)
        current = sorted(p.name for p in draft.parent.glob("*.dxf"))
        self.assertTrue(any(name.endswith("_Изм1.dxf") and "усиленная" not in name for name in current),
                        f"развёртка новой ревизии БЧ-детали: {current}")
        self.assertFalse(draft.exists(), "прежняя развёртка ушла из папки выдачи")
        archived = sorted(p.name for p in (draft.parent / "_Аннулировано").glob("*.dxf"))
        self.assertTrue(any(name.startswith(draft.stem) for name in archived), f"прежняя развёртка в архиве: {archived}")
        self.assertTrue(decoy.exists(), "файл другого документа («… усиленная») не тронут")
        exported = (product / "_Экспорт.txt").read_text(encoding="utf-8-sig")
        self.path("export.txt").write_text(exported, encoding="utf-8")
        self.assertNotIn("оформите новую ревизию", exported, "БЧ-деталь с ревизией не пропущена как выданная")
        # Прежняя развёртка унесена в архив уже после записи отчёта выгрузки — ссылки на неё в отчёте нет (ревью 23.09.2026).
        self.assertNotIn(draft.name, exported, "отчёт выгрузки не ссылается на файл из «_Аннулировано»")
        self.assertTrue(any(name.endswith("_Изм1.dxf") and name in exported for name in current), exported)
        self.s.close_all()
        self.assertEqual("1", oracles.value(self.persisted(part), "Ревизия"), "ревизия БЧ — в её «Ревизии»")
        self.assertEqual([], self.addin_errors(), "ошибки в журнале надстройки")

    def test_V06_unsaved_stamp_is_rolled_back(self):
        """V06 (аудит 23.09.2026, SAVE-9): штамп ревизии не сохранился (файл чертежа стал только для чтения уже после
        открытия). Раньше строка журнала оставалась, а новый номер — в открытом чертеже: следующее сохранение записало
        бы ревизию без выгрузки. Теперь штамп в открытом чертеже прежний, строка журнала помечена «ОТМЕНЕНО», файлы
        выдачи не тронуты."""
        product, models = self._product()
        drawing = models / DRAWING
        doc = self.s.open(drawing)
        self.s.activate(doc)
        cpm = com.dyn(doc).Extension.CustomPropertyManager("")
        before = com.prop_get(cpm, "Revision")
        os.chmod(drawing, stat.S_IREAD)
        try:
            status = self._revision()
        finally:
            os.chmod(drawing, stat.S_IREAD | stat.S_IWRITE)
        self.path("status.txt").write_text(status, encoding="utf-8")
        self.assertTrue(status.startswith("error|"), status)
        self.assertIn("не сохранён", status, status)
        self.assertIn("ОТМЕНЕНО", status, status)
        self.assertEqual(before, com.prop_get(com.dyn(doc).Extension.CustomPropertyManager(""), "Revision"),
                         "свойство ревизии в открытом чертеже — прежнее")
        rows = list(openpyxl.load_workbook(product / "Изменения.xlsx")["Изменения"].iter_rows(values_only=True))
        self.assertEqual(2, len(rows), f"строка журнала не удалена: {rows}")
        self.assertTrue(str(rows[1][5]).startswith("ОТМЕНЕНО: "), f"строка помечена отменённой: {rows[1]}")
        self.assertFalse((product / "02_PDF").exists() and list((product / "02_PDF").glob("*_Изм1.pdf")),
                         "выгрузки новой ревизии нет")
        self.s.close_all()

    def test_V07_check_flags_issued_document_changed_without_revision(self):
        """V07 (аудит 23.09.2026, CHK-13): изделие выдано (`_Выдано_…` с суммами), потом деталь поправили и сохранили без
        новой ревизии — цех работает по старому, а правка до него не дойдёт. Проверка изделия даёт замечание и называет
        чертёж, на котором поднять ревизию; после «Новой ревизии» замечания нет."""
        product, models = self._product(issued=False)
        lines = ["Готово к производству", "Изделие:  ПРТИ.468211.100", "Отметил:  Тестов Т.Т., 10.09.2026 12:00",
                 "Журнал:   строка 0 (Изменения.xlsx)", "", "Документы изделия:"]
        for f in sorted(models.iterdir()):
            lines.append("  " + hashlib.sha256(f.read_bytes()).hexdigest() + "  " + f.name)
        (product / "_Выдано_2026-09-10_1200.txt").write_text("\n".join(lines) + "\n", encoding="utf-8-sig")

        part = self.s.open(models / PART)
        com.prop_set(com.dyn(part).Extension.CustomPropertyManager(""), "Примечание", "правка после выдачи")
        com.dyn(part).SetSaveFlag()
        self.s.save(part)
        self.s.close_all()

        def check():
            asm = self.s.open(models / ASM)
            self.s.activate(asm)
            com.call(self.s.eskd(), "CheckProductSilent")
            status = str(com.call(self.s.eskd(), "CheckStatus"))
            self.assertTrue(status.startswith("ok|"), status)
            return (product / "_Проверка.txt").read_text(encoding="utf-8-sig")

        report = check()
        self.path("check1.txt").write_text(report, encoding="utf-8")
        self.assertEqual(1, report.count("изменён после выдачи"), report)
        self.assertIn(PART + " — изменён после выдачи в производство (_Выдано_2026-09-10_1200.txt)", report, report)
        self.assertIn("«Новая ревизия» в «" + DRAWING + "»", report, report)

        drawing = self.s.open(models / DRAWING)
        self.s.activate(drawing)
        status = self._revision()
        self.assertTrue(status.startswith("ok|"), status)
        self.s.close_all()
        report = check()
        self.path("check2.txt").write_text(report, encoding="utf-8")
        self.assertNotIn("изменён после выдачи", report, "ревизия оформлена — замечания нет")
        self.s.close_all()

    def test_V03_draft_document_needs_no_revision(self):
        """V03: изделие не выдано — кнопка серая, вызов отказывает и ничего не пишет (Т-48)."""
        product, models = self._product(issued=False)
        doc = self.s.open(models / DRAWING)
        self.s.activate(doc)
        self.assertEqual(0, int(com.call(self.s.eskd(), "EnableRevisionCommand")), "кнопка серая у черновика")
        self.assertIn("не выдан", str(com.call(self.s.eskd(), "RevisionUnavailable")), "подсказка объясняет почему")
        status = self._revision()
        self.assertTrue(status.startswith("error|"), status)
        self.assertFalse((product / "Изменения.xlsx").exists(), "журнал не заведён")

    def test_V04_assembly_has_no_revision(self):
        """V04: у сборки ревизии нет — ревизия принадлежит чертежу (Р0-8)."""
        product, models = self._product()
        doc = self.s.open(models / ASM)
        self.s.activate(doc)
        self.assertEqual(0, int(com.call(self.s.eskd(), "EnableRevisionCommand")), "кнопка серая у сборки")
        self.assertTrue(self._revision().startswith("error|"), "вызов отказал")
        self.assertNoPropertyWrites()


if __name__ == "__main__":
    unittest.main()
