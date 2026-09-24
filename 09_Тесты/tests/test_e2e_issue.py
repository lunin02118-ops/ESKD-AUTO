# -*- coding: utf-8 -*-
"""E2E, группы R и Y — кнопки «Готово к производству» (ТЗ-04 Р4-8) и «Закрыть заказ» (ТЗ-02 Т-44…Т-47).

Окна «Выдать в производство» на заказ больше нет: изделие отмечается готовым по одному, отчёт _Выдано лежит
в папке изделия, PDF листов участков делает Excel из книги ЛЗК. Книги здесь собираются вручную (openpyxl) —
кнопка читает их как файлы.
"""
import os
import shutil
import time
import unittest
from pathlib import Path

import openpyxl

from eskd_e2e import com, paths
from eskd_e2e.testing import SwTestCase

PRODUCT = "И01_ПРТИ.468211.100"
CIPHER = "ПРТИ.468211.100"
ASM = "ПРТИ.468211.100 СБ Кондуктор сварочный.sldasm"
ORDER = "2026-010 Стенд"
DOCS = "04_Сопроводительная документация"


class Issue(SwTestCase):

    # ------------------------------------------------------------------ фикстура заказа
    def _order(self, issued=False):
        """Заказ с одним изделием; возвращает (заказ, изделие, сборка)."""
        short = self._case_name().split("_")[1]
        subdir = f"{short}/_Заявки/{ORDER}/02_Металл/{PRODUCT}/01_3D"
        models = self.s.run_dir / subdir
        if models.exists():
            shutil.rmtree(models, ignore_errors=True)
        for src in sorted(Path(paths.FIXTURES_A).iterdir()):
            if src.suffix.lower() in (".sldprt", ".sldasm", ".slddrw"):
                self.s.workspace_copy(src, subdir=subdir)
        product = models.parent
        order = product.parent.parent
        if issued:
            (product / "_Выдано_2026-09-10_0900.txt").write_text(
                "Готово к производству\nИзделие: " + CIPHER + "\n", encoding="utf-8")
        return order, product, models / ASM

    def _book(self, path):
        """Книга ЛЗК нового образца: паспорт с именем «Тираж» и лист участка."""
        from openpyxl.workbook.defined_name import DefinedName

        path.parent.mkdir(parents=True, exist_ok=True)
        book = openpyxl.Workbook()
        passport = book.active
        passport.title = "Паспорт"
        passport["A9"] = "Изделий в заказе, шт."
        passport["B9"] = 1
        book.defined_names.add(DefinedName("Тираж", attr_text="'Паспорт'!$B$9"))
        blank = book.create_sheet("Заготовительный")
        blank["A1"] = "Заготовительный участок"
        blank["A6"] = 1
        book.save(path)
        return path

    def _ready(self, asm):
        doc = self.s.open(asm)
        self.s.activate(doc)
        com.call(self.s.eskd(), "ReadyForProductionSilent")
        return str(com.call(self.s.eskd(), "ReadyStatus"))

    def _close(self, order, archive):
        com.call(self.s.eskd(), "CloseOrderSilent", str(order), str(archive))
        return str(com.call(self.s.eskd(), "CloseOrderStatus"))

    def _nothing_issued(self, product):
        self.assertEqual([], list(product.glob("_Выдано_*.txt")), "отметка «готово» не написана")
        self.assertEqual([], list((product / "02_PDF").glob("ЛЗК_*.pdf")) if (product / "02_PDF").exists() else [],
                         "PDF листов участков не сделаны")

    # ------------------------------------------------------------------ R: готово к производству
    def test_G01_without_book_asks_for_lzk(self):
        """G01: у изделия нет книги ЛЗК — кнопка просит сначала «Ведомость ЛЗК» и ничего не пишет."""
        _, product, asm = self._order()
        status = self._ready(asm)
        self.assertTrue(status.startswith("error|"), status)
        self.assertIn("Ведомость ЛЗК", status, status)
        self._nothing_issued(product)

    def test_G02_legacy_book_must_be_rebuilt(self):
        """G02: ведомость старого образца (в корне изделия, без участков) — кнопка просит пересобрать книгу."""
        _, product, asm = self._order()
        openpyxl.Workbook().save(product / f"Ведомость_{CIPHER}.xlsx")
        status = self._ready(asm)
        self.assertTrue(status.startswith("error|"), status)
        self.assertIn("старого образца", status, status)
        self._nothing_issued(product)

    def test_G03_unchecked_product_is_not_ready(self):
        """G03: книга есть, но проверка изделия не «ГОТОВО» — отказ с причиной, PDF и отметки нет."""
        _, product, asm = self._order()
        self._book(product / DOCS / f"ЛЗК_{CIPHER}.xlsx")
        status = self._ready(asm)
        self.assertTrue(status.startswith("error|"), status)
        self.assertIn("ГОТОВО", status, f"названа причина отказа: {status}")
        self._nothing_issued(product)
        self.assertEqual([], self.addin_errors(), "ошибки в журнале надстройки")

    def test_G04_product_ready_report_protects_its_documents(self):
        """G04: отметка «готово» лежит в папке изделия — выгрузка видит документ выданным (Т-30)."""
        _, product, asm = self._order()
        (product / "_Выдано_2026-09-11_1200.txt").write_text(
            "\n".join(["Готово к производству", "", "Документы изделия:", "  ab12  " + ASM, ""]), encoding="utf-8")
        doc = self.s.open(asm)
        self.s.activate(doc)
        com.call(self.s.eskd(), "ExportProductSilent")
        status = str(com.call(self.s.eskd(), "ExportStatus"))
        self.assertTrue(status.startswith("ok|"), status)
        text = Path(status.split("|")[3]).read_text(encoding="utf-8-sig")
        self.assertIn("документ выдан в производство, оформите новую ревизию", text,
                      "отметка в папке изделия защищает выданный документ")

    def _wait(self, start, status_method, running="running", timeout=900):
        com.call(self.s.eskd(), start)
        deadline = time.time() + timeout
        status = running
        while time.time() < deadline:
            status = str(com.call(self.s.eskd(), status_method))
            if status and status != running:
                break
            time.sleep(1)
        return status

    def test_G05_checked_product_becomes_ready(self):
        """G05: изделие с чертежами, выгрузкой и книгой ЛЗК проходит проверку «ГОТОВО» — кнопка делает PDF листов
        участков через Excel, пишет отметку «готово» и отчёт _Выдано в папке изделия."""
        order, product, asm = self._order()
        # Изделие доводится до готового кнопками конструктора: «Синхронизировать» на каждой модели,
        # «Деталь БЧ» — у деталей без чертежа; модели сохраняются, как сохранил бы их конструктор.
        for model in sorted((product / "01_3D").iterdir()):
            # покупные (без обозначения) приходят готовыми — их кнопки не трогают
            if model.suffix.lower() not in (".sldprt", ".sldasm") or model.name == ASM or not model.name.startswith("ПРТИ"):
                continue
            doc = self.s.open(model)
            self.s.activate(doc)
            self.assertGreaterEqual(int(com.call(self.s.eskd(), "SyncActiveDocumentSilent")), 0, model.name)
            # кнопка «Деталь БЧ» вдавлена (3) у уже оформленной — повторное нажатие сняло бы оформление
            bch = int(com.call(self.s.eskd(), "EnableBchCommand"))
            if model.suffix.lower() == ".sldprt" and not model.with_suffix(".SLDDRW").exists() and bch != 3:
                self.assertEqual(1, int(com.call(self.s.eskd(), "ToggleDrawinglessSilent")), model.name)
            # Рама сварная фикстуры — три тела в одной детали: в IGS она не идёт, и с «Лазерная резка трубы» проверка
            # изделия её не пропустит (З-51, её проверяет K10). Здесь конструктор решил резать раму не на труборезе.
            if model.name == "ПРТИ.468211.105 Рама сварная.sldprt":
                from eskd_e2e import build
                build.props(doc, {"Операции": "Покраска"})
            self.s.save(doc)
            self.s.close(doc)
        norms = paths.ROOT / "02_Шаблоны_и_Форматки" / "Справочники" / "Нормативы_производства.xlsx"
        (order / norms.name).write_bytes(norms.read_bytes())
        doc = self.s.open(asm)
        self.s.activate(doc)
        # Порядок работы (решение владельца 23.09.2026, З-27): проверка → ЛЗК → выгрузка → «Готово». Книга и выгрузка
        # помнят версию изделия, по которой сделаны; «Готово» требует, чтобы она была текущей.
        # Первая проверка — как у конструктора, «Применить и сохранить»: она дописывает реквизиты сборки. Сохранения ЛЗК
        # и выгрузки синхронизацию при сохранении не запускают.
        self.assertTrue(self._wait("CheckProductApplySilent", "CheckStatus", running="").startswith("ok|"), "проверка")
        lzk = self._wait("BuildLzkSilent", "LzkStatus")
        self.assertTrue(lzk.startswith("ok|"), lzk)
        self.assertTrue(self._wait("ExportProductSilent", "ExportStatus", running="").startswith("ok|"), "выгрузка")
        check = self._wait("CheckProductSilent", "CheckStatus", running="")
        report = product / "_Проверка.txt"
        self.assertTrue(check.startswith("ok|ГОТОВО|"),
                        check + "\n" + (report.read_text(encoding="utf-8-sig") if report.exists() else ""))

        com.call(self.s.eskd(), "ReadyForProductionSilent")
        status = str(com.call(self.s.eskd(), "ReadyStatus"))
        self.assertTrue(status.startswith("ok|"), status)
        _, issued, sheets, folder = status.split("|")
        self.assertEqual(str(product).lower(), folder.lower(), "папка изделия в итоге")
        pdfs = sorted((product / "02_PDF").glob(f"ЛЗК_{CIPHER}_*.pdf"))
        self.assertEqual(int(sheets), len(pdfs), f"PDF на каждый лист участка и «Расход»: {[p.name for p in pdfs]}")
        self.assertIn(f"ЛЗК_{CIPHER}_Расход.pdf", [p.name for p in pdfs], "«Расход» в цех")
        for pdf in pdfs:
            self.assertEqual(b"%PDF", pdf.read_bytes()[:4], f"{pdf.name} — PDF от Excel")
        issued = Path(issued)
        self.assertEqual(str(product).lower(), str(issued.parent).lower(), "отчёт _Выдано в папке изделия")
        self.assertTrue(issued.name.startswith("_Выдано_"), issued.name)
        text = issued.read_text(encoding="utf-8-sig")
        for pdf in pdfs:
            self.assertIn(pdf.name, text, f"{pdf.name} в отчёте")
        self.assertIn(f"ЛЗК_{CIPHER}.xlsx", text, "книга в отчёте")
        self.assertIn(ASM.lower(), text.lower(), "документы изделия перечислены — их защищает Т-30")
        book = openpyxl.load_workbook(product / DOCS / f"ЛЗК_{CIPHER}.xlsx")
        ws, cell = next(iter(book.defined_names["Выдано"].destinations))
        self.assertTrue(str(book[ws][cell.replace("$", "")].value or "").strip(), "отметка «Выдано» в паспорте книги")
        self.assertEqual([], self.addin_errors(), "ошибки в журнале надстройки")

    # ------------------------------------------------------------------ Y: закрытие заказа
    def test_Y01_close_packs_order_and_moves_it_to_archive(self):
        """Y01: комплекты Pack and Go, сверка, архив Y:\\<год>\\<заказ> и папка в _Сдано с _Архив.txt (Т-45…Т-47).
        Сборка комплекта открывается из архива сама: все её компоненты лежат рядом. 24.09.2026 Pack and Go
        SolidWorks 2025 SP3 перечислял только сборку и чертежи, и комплект уезжал в архив без единой детали."""
        order, product, asm = self._order(issued=True)
        archive = self.s.run_dir / self._case_name() / "archive"
        archive.mkdir(parents=True, exist_ok=True)
        self.s.close_all()

        status = self._close(order, archive)
        self.assertTrue(status.startswith("ok|"), status)
        _, target, files, done = status.split("|")
        self.assertGreater(int(files), 0, status)

        target = Path(target)
        self.assertTrue(target.is_dir(), f"папка архива: {target}")
        self.assertEqual(archive.name, target.parent.parent.name, "архив в указанном корне")
        kit = target / "Комплекты" / CIPHER
        self.assertTrue(kit.is_dir(), f"комплект изделия: {sorted(p.name for p in target.iterdir())}")
        packed = sorted(p.name for p in kit.iterdir())
        self.assertTrue(any(p.lower().endswith(".sldasm") for p in packed), f"сборка в комплекте: {packed}")
        self.assertTrue(any(p.lower().endswith(".slddrw") for p in packed), f"чертежи в комплекте: {packed}")
        self.assertTrue((target / "02_Металл" / PRODUCT / "01_3D" / ASM).is_file(), "файлы заказа скопированы")
        # Папка заказа уже в «_Сдано»: записанный в сборке путь компонентов не существует, и SolidWorks найдёт их
        # только рядом со сборкой комплекта.
        top = self.s.open(kit / ASM, readonly=True)
        try:
            deps = [str(p) for p in com.as_list(com.call(top, "GetDependencies2", True, True, False))[1::2]]
        finally:
            self.s.close(top)
        self.assertGreaterEqual(len(deps), 8, f"компоненты сборки комплекта: {deps}")
        outside = [p for p in deps if Path(p).parent != kit]
        self.assertEqual([], outside, f"сборка комплекта берёт компоненты не из комплекта; в комплекте: {packed}")

        done = Path(done)
        self.assertFalse(order.exists(), "папка заказа убрана из рабочего каталога")
        self.assertTrue(done.is_dir(), f"заказ в «_Сдано»: {done}")
        note = (done / "_Архив.txt").read_text(encoding="utf-8-sig")
        self.assertIn(str(target), note, f"записка ссылается на архив: {note}")
        self.assertFalse((order.parent / "_tmp" / ORDER).exists(), "временная папка убрана")
        self.assertEqual([], self.addin_errors(), "ошибки в журнале надстройки")

    def test_Y02_open_document_and_missing_issue_report_stop_closing(self):
        """Y02: заказ не выдавался или его документ открыт — отказ без единого движения файлов (Т-44)."""
        order, product, asm = self._order()
        archive = self.s.run_dir / self._case_name() / "archive"
        archive.mkdir(parents=True, exist_ok=True)

        status = self._close(order, archive)
        self.assertTrue(status.startswith("error|"), status)
        self.assertIn("не выдавался", status, status)

        (product / "_Выдано_2026-09-10_0900.txt").write_text("Готово к производству\n", encoding="utf-8")
        doc = self.s.open(asm)
        self.s.activate(doc)
        status = self._close(order, archive)
        self.assertTrue(status.startswith("error|"), status)
        self.assertIn("открыт", status, status)
        self.assertTrue(order.is_dir(), "папка заказа на месте")
        self.assertEqual([], list(archive.iterdir()), "в архиве пусто")

    def test_Y06_namesake_from_other_order_stops_closing(self):
        """Y06 (сверка SW API 23.09.2026, №16): в SolidWorks открыта сборка другого заказа с теми же именами файлов.
        SolidWorks берёт компонент сначала из открытых документов — по имени, а не по пути, — и Pack and Go заказа А
        упаковал бы в архив деталь заказа Б. Теперь «Закрыть заказ» отказывает до любых действий с файлами; когда
        чужие документы закрыты, в архив уходит своя деталь."""
        from eskd_e2e import build, oracles

        order, product, asm = self._order(issued=True)
        archive = self.s.run_dir / self._case_name() / "archive"
        archive.mkdir(parents=True, exist_ok=True)
        short = self._case_name().split("_")[1]
        other = f"{short}/_Заявки/2026-011 Другой/02_Металл/{PRODUCT}/01_3D"
        for src in sorted(Path(paths.FIXTURES_A).iterdir()):
            if src.suffix.lower() in (".sldprt", ".sldasm"):
                self.s.workspace_copy(src, subdir=other)
        stand = "ПРТИ.468211.102 Стойка.sldprt"
        # Сначала деталь: открытая сборка уже держит её в памяти, и второй раз её не открыть.
        part = self.s.open(self.s.run_dir / other / stand)
        build.props(part, {"Заказ": "Б"})
        self.s.save(part)
        foreign = self.s.open(self.s.run_dir / other / ASM)
        self.s.activate(foreign)

        status = self._close(order, archive)
        self.assertTrue(status.startswith("error|"), status)
        self.assertIn("одноимённ", status, status)
        # Путь документа SolidWorks отдаёт в регистре файла на диске («.SLDPRT»), а не как его назвали в тесте.
        self.assertIn(stand.lower(), status.lower(), status)
        self.assertTrue(order.is_dir(), "папка заказа на месте")
        self.assertEqual([], list(archive.iterdir()), "в архиве пусто")
        self.assertFalse((order.parent / "_tmp" / ORDER).exists(), "временная папка не заведена")

        self.s.close_all()
        status = self._close(order, archive)
        self.assertTrue(status.startswith("ok|"), status)
        kit = Path(status.split("|")[1]) / "Комплекты" / CIPHER
        self.assertIsNone(oracles.value(oracles.read_persisted(self.s, kit / stand), "Заказ"), "в архив ушла своя Стойка")
        self.assertEqual([], self.addin_errors(), "ошибки в журнале надстройки")

    def test_Y07_namesake_from_other_order_stops_check_export_and_lzk(self):
        """Y07 (№16, решение владельца 24.09.2026 — «как правильно»): открыта деталь другого заказа с тем же именем, и
        сборка заказа А взяла её вместо своей. Раньше проверка писала общее «ссылка на другой заказ», выгрузка молча
        пропускала «чужую» деталь, а ЛЗК брала в книгу её размеры. Теперь проверка называет подмену и что сделать,
        выгрузка и ЛЗК останавливаются до записи файлов."""
        from eskd_e2e import build

        _, product, asm = self._order(issued=False)
        short = self._case_name().split("_")[1]
        other = f"{short}/_Заявки/2026-011 Другой/02_Металл/{PRODUCT}/01_3D"
        for src in sorted(Path(paths.FIXTURES_A).iterdir()):
            if src.suffix.lower() in (".sldprt", ".sldasm"):
                self.s.workspace_copy(src, subdir=other)
        stand = "ПРТИ.468211.102 Стойка.sldprt"
        part = self.s.open(self.s.run_dir / other / stand)
        build.props(part, {"Заказ": "Б"})
        self.s.save(part)
        doc = self.s.open(asm)
        self.s.activate(doc)
        comps = [str(com.call(c, "GetPathName")) for c in com.as_list(com.dyn(doc).GetComponents(False)) or []]
        self.assertTrue(any(p.lower() == str(self.s.run_dir / other / stand).lower() for p in comps),
                        f"предусловие: сборка А взяла Стойку заказа Б: {comps}")

        com.call(self.s.eskd(), "CheckProductSilent")
        text = (product / "_Проверка.txt").read_text(encoding="utf-8-sig")
        self.assertIn("одноимённый файл из другой папки", text, text[-2000:])
        # Имена — в регистре файла на диске («.SLDPRT»).
        self.assertIn(stand.lower(), text.lower(), text[-2000:])

        before = {p for p in product.rglob("*") if p.is_file()}
        com.call(self.s.eskd(), "ExportProductSilent")
        status = str(com.call(self.s.eskd(), "ExportStatus"))
        self.assertTrue(status.startswith("error|") and "одноимённ" in status, status)
        self.assertIn(stand.lower(), status.lower(), status)

        com.call(self.s.eskd(), "BuildLzkSilent")
        deadline = time.time() + 120
        status = "running"
        while time.time() < deadline and status == "running":
            time.sleep(1)
            status = str(com.call(self.s.eskd(), "LzkStatus"))
        self.assertTrue(status.startswith("error|") and "одноимённ" in status, status)
        after = {p for p in product.rglob("*") if p.is_file()}
        self.assertEqual(sorted(str(p) for p in after - before if p.name != "_Проверка.txt"), [],
                         "выгрузка и ЛЗК файлов не писали")
        self.assertEqual([], self.addin_errors(), "ошибки в журнале надстройки")
        self.s.close_all()

    def test_Y03_unavailable_archive_keeps_order_untouched(self):
        """Y03: архивный диск недоступен — заказ не тронут, объяснение понятное (Т-46)."""
        order, _, _ = self._order(issued=True)
        self.s.close_all()
        status = self._close(order, self.s.run_dir / self._case_name() / "нет-такого-диска")
        self.assertTrue(status.startswith("error|"), status)
        self.assertIn("недоступен", status, status)
        self.assertTrue((order / "02_Металл" / PRODUCT / "01_3D" / ASM).is_file(), "файлы заказа на месте")
        self.assertFalse((order.parent / "_Сдано" / ORDER).exists(), "заказ не сдан")

    def _close_to_default_archive(self, order):
        """Закрыть заказ без адреса архива: архив по умолчанию — «_Архив» рядом с корнем заказов (19.09.2026)."""
        orders_root = order.parent
        archive = orders_root.parent / "_Архив"
        self.s.close_all()
        status = self._close(order, "")
        self.assertTrue(status.startswith("ok|"), status)
        _, target, files, done = status.split("|")
        target, done = Path(target), Path(done)
        self.assertGreater(int(files), 0, status)
        self.assertEqual(str(archive / time.strftime("%Y") / order.name).lower(), str(target).lower(), "архив _Архив\\<год>\\<заказ>")
        self.assertTrue((target / "Комплекты" / CIPHER).is_dir(), f"комплект изделия: {sorted(p.name for p in target.iterdir())}")
        self.assertTrue((target / "02_Металл" / PRODUCT / "01_3D" / ASM).is_file(), "файлы заказа в архиве")
        self.assertEqual(str(orders_root / "_Сдано" / order.name).lower(), str(done).lower(), "заказ в «_Сдано»")
        self.assertFalse(order.exists(), "папка заказа убрана из рабочего каталога")
        self.assertIn(str(target), (done / "_Архив.txt").read_text(encoding="utf-8-sig"), "записка ссылается на архив")
        self.assertFalse((orders_root / "_tmp" / order.name).exists(), "временная папка убрана")
        self.assertEqual([], self.addin_errors(), "ошибки в журнале надстройки")
        return target, done

    def test_Y04_default_archive_is_next_to_orders_root(self):
        """Y04: диска Y: нет — «Закрыть заказ» без адреса кладёт архив в «_Архив» рядом с «_Заявки»."""
        order, _, _ = self._order(issued=True)
        (order.parent.parent / "_Архив").mkdir(exist_ok=True)
        self._close_to_default_archive(order)

    @unittest.skipUnless(os.environ.get("ESKD_LIVE_ORDERS_ROOT"), "живая проверка NAS: задать ESKD_LIVE_ORDERS_ROOT=<…\\_Заявки>")
    def test_Y05_live_close_on_nas(self):
        """Y05 (по запросу): тестовый заказ закрывается на настоящем NAS в настоящий «_Архив»; убирается только он сам."""
        orders_root = Path(os.environ["ESKD_LIVE_ORDERS_ROOT"])
        archive = orders_root.parent / "_Архив"
        self.assertTrue(archive.is_dir(), f"нет папки архива {archive}")
        name = "_тест_ЕСКД_закрытие_" + time.strftime("%Y%m%d_%H%M%S")
        order = orders_root / name
        models = order / "02_Металл" / PRODUCT / "01_3D"
        models.mkdir(parents=True)
        tmp_existed = (orders_root / "_tmp").exists()
        target = done = None
        try:
            for src in sorted(Path(paths.FIXTURES_A).iterdir()):
                if src.suffix.lower() in (".sldprt", ".sldasm", ".slddrw"):
                    shutil.copy2(src, models / src.name)
            (models.parent / "_Выдано_2026-09-10_0900.txt").write_text(
                "Готово к производству\nИзделие: " + CIPHER + "\n", encoding="utf-8")
            target, done = self._close_to_default_archive(order)
        finally:
            # Только то, что создал сам тест: тестовый заказ, его архив и его папка в «_Сдано».
            for path in (order, done, target, orders_root / "_tmp" / name):
                if path is not None and Path(path).name == name and Path(path).exists():
                    shutil.rmtree(path, ignore_errors=True)
            year = archive / time.strftime("%Y")
            if year.is_dir() and not any(year.iterdir()):
                year.rmdir()
            tmp = orders_root / "_tmp"
            if not tmp_existed and tmp.is_dir() and not any(tmp.iterdir()):
                tmp.rmdir()


if __name__ == "__main__":
    unittest.main()
