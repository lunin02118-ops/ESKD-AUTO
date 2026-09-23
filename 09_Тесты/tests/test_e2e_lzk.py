# -*- coding: utf-8 -*-
"""E2E, группа L — кнопка «Ведомость ЛЗК» (ТЗ-02 Т-35…Т-37, ТЗ-04): операции, габарит, выгрузка SWTools без окна,
живая книга ЛЗК — паспорт, участки, расход, нормы — в «04_Сопроводительная документация» изделия."""
import os
import shutil
import time
import unittest
import winreg
from pathlib import Path

from eskd_e2e import com, oracles, paths
from eskd_e2e.testing import SwTestCase, tags

V = oracles.value
PRODUCT = "И01_ПРТИ.468211.100_Кондуктор"
ASM = "ПРТИ.468211.100 СБ Кондуктор сварочный.sldasm"
SHEET_PART = "ПРТИ.468211.101 Пластина опорная.sldprt"
SUBASSEMBLY = "ПРТИ.468211.110 СБ Узел опоры.sldasm"
DOCS = "04_Сопроводительная документация"
BOOK = "ЛЗК_ПРТИ.468211.100.xlsx"
SHEETS = ["Паспорт", "Ведомость", "Сводная", "Заготовительный", "Сварочный", "Покрасочный", "Комплектовочный", "Расход", "Нормы"]
TIMEOUT = 600


def swtools_installed():
    for view in (winreg.KEY_WOW64_64KEY, winreg.KEY_WOW64_32KEY):
        try:
            with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\SWTools\InstallRoot", 0,
                                winreg.KEY_READ | view) as key:
                root = winreg.QueryValueEx(key, "")[0]
                if (Path(root) / "SWTools.exe").is_file():
                    return True
        except OSError:
            continue
    return False


@unittest.skipUnless(swtools_installed(), "SWTools не установлен")
class Lzk(SwTestCase):

    def _product(self):
        """Копия изделия в структуре заказа: И01_<шифр>_<имя>\\01_3D со всеми моделями фикстуры A."""
        models = self.case_dir / PRODUCT / "01_3D"
        models.mkdir(parents=True, exist_ok=True)
        # справочник нормативов «на заказ» над изделием: встроенных нормативов у надстройки нет (Т-13)
        # справочник бланков — тоже из репозитория: иначе надстройка возьмёт опубликованный в инструментарии на NAS
        for name in ("Нормативы_производства.xlsx", "ЛЗК_бланки.xlsx"):
            ref = paths.ROOT / "02_Шаблоны_и_Форматки" / "Справочники" / name
            (self.case_dir / ref.name).write_bytes(ref.read_bytes())
        for src in sorted(Path(paths.FIXTURES_A).iterdir()):
            if src.suffix.lower() in (".sldprt", ".sldasm"):
                self.s.workspace_copy(src, subdir=f"{self._case_name()}/{PRODUCT}/01_3D")
        return self.case_dir / PRODUCT, models / ASM

    def _build(self):
        com.call(self.s.eskd(), "BuildLzkSilent")
        deadline = time.time() + TIMEOUT
        status = "running"
        while time.time() < deadline:
            status = str(com.call(self.s.eskd(), "LzkStatus"))
            if status != "running":
                break
            time.sleep(1)
        return status

    def _outside(self, folder):
        """Сборка в папке изделия, её детали — в папке folder вне изделия (относительно каталога теста). Возвращает
        (папка изделия, пути деталей, итог кнопки, замечания)."""
        from eskd_e2e import build

        product = self.case_dir / PRODUCT
        (product / "01_3D").mkdir(parents=True, exist_ok=True)
        for name in ("Нормативы_производства.xlsx", "ЛЗК_бланки.xlsx"):
            ref = paths.ROOT / "02_Шаблоны_и_Форматки" / "Справочники" / name
            (self.case_dir / ref.name).write_bytes(ref.read_bytes())
        parts = [self.s.workspace_copy(Path(paths.FIXTURES_A) / n, subdir=f"{self._case_name()}/{folder}")
                 for n in (SHEET_PART, "ПРТИ.468211.111 Стойка трубная.sldprt")]
        asm, opened = build.assembly(self.s, [(parts[0], 0, 0, 0), (parts[1], 0, 0.15, 0)])
        self.s.save_as(asm, product / "01_3D" / ASM)
        for doc in opened:
            self.s.close(doc)
        self.s.activate(asm)
        status = self._build()
        notices = str(com.call(self.s.eskd(), "LastNotices"))
        self.path("outcome.txt").write_text(status + "\n\n" + notices, encoding="utf-8")
        self.assertTrue(status.startswith("ok|"), f"{status}\n{notices}")
        return product, parts, status, notices

    def _sections(self, workbook):
        """{обозначение: операции} ведомости и обозначения на листах участков."""
        import openpyxl
        wb = openpyxl.load_workbook(workbook)
        main = wb["Ведомость"]
        ops = {main.cell(r, 3).value: main.cell(r, 9).value for r in range(7, main.max_row + 1) if main.cell(r, 10).value}
        on_sections = {wb[s].cell(r, 2).value for s in ("Заготовительный", "Сварочный", "Покрасочный", "Комплектовочный")
                       for r in range(6, wb[s].max_row + 1)}
        return ops, on_sections

    def test_L04_parts_outside_product_folder(self):
        """L04 (замечание владельца 19.09.2026): сборка в папке изделия, её детали — в другой папке. Книга ЛЗК — на всю
        сборку: у деталей из другой папки операции (не «?»), записанные в их модели, и они есть на участках."""
        product, parts, _, notices = self._outside("Детали_общие")
        workbook = product / DOCS / BOOK
        self.assertTrue(workbook.is_file(), "книга ЛЗК изделия")
        ops, on_sections = self._sections(workbook)
        for designation, path in (("ПРТИ.468211.101", parts[0]), ("ПРТИ.468211.111", parts[1])):
            with self.subTest(part=designation):
                self.assertNotIn(ops.get(designation) or "?", ("?", ""), f"операции детали из другой папки: {ops}\n{notices}")
                self.assertIn(designation, on_sections, f"деталь из другой папки на участках: {on_sections}")
        self.s.close_all()
        for designation, path in (("ПРТИ.468211.101", parts[0]), ("ПРТИ.468211.111", parts[1])):
            with self.subTest(disk=designation):
                self.assertEqual(ops[designation], V(self.persisted(path), "Операции"), "операции записаны в модель")

    def test_L12_other_order_models_get_operations_only_in_book(self):
        """L12 (решение владельца 23.09.2026): книга ЛЗК одного заказа не меняет модели другого. Деталь другого заказа
        получает операции в книге, её файл не меняется и в SolidWorks не помечен изменённым; деталь из другой папки того
        же заказа записывается (без окна — как с галочкой «В модель», которая стоит по умолчанию)."""
        from eskd_e2e import build

        # Каталог назван коротко: с полным именем теста путь к сборке переваливает за 260 знаков, и SolidWorks молча
        # не сохраняет её.
        case = "L12"
        root = self.s.run_dir / case
        shutil.rmtree(root, ignore_errors=True)
        product = root / "03_ЗАКАЗЫ" / "778_Тест" / "02_Металл" / PRODUCT
        (product / "01_3D").mkdir(parents=True, exist_ok=True)
        for name in ("Нормативы_производства.xlsx", "ЛЗК_бланки.xlsx"):
            ref = paths.ROOT / "02_Шаблоны_и_Форматки" / "Справочники" / name
            (root / ref.name).write_bytes(ref.read_bytes())
        foreign = self.s.workspace_copy(Path(paths.FIXTURES_A) / SHEET_PART,
                                        subdir=f"{case}/03_ЗАКАЗЫ/775_Другой/02_Металл/И01_775_Стол/01_3D")
        same = self.s.workspace_copy(Path(paths.FIXTURES_A) / "ПРТИ.468211.111 Стойка трубная.sldprt",
                                     subdir=f"{case}/03_ЗАКАЗЫ/778_Тест/Общие детали")
        before = foreign.read_bytes()
        asm, opened = build.assembly(self.s, [(foreign, 0, 0, 0), (same, 0, 0.15, 0)])
        self.s.save_as(asm, product / "01_3D" / ASM)
        for doc in opened:
            self.s.close(doc)
        self.s.activate(asm)
        status = self._build()
        notices = str(com.call(self.s.eskd(), "LastNotices"))
        self.path("outcome.txt").write_text(status + "\n\n" + notices, encoding="utf-8")
        self.assertTrue(status.startswith("ok|"), f"{status}\n{notices}")
        ops, on_sections = self._sections(product / DOCS / BOOK)
        for designation in ("ПРТИ.468211.101", "ПРТИ.468211.111"):
            with self.subTest(book=designation):
                self.assertNotIn(ops.get(designation) or "?", ("?", ""), f"операции в книге: {ops}\n{notices}")
                self.assertIn(designation, on_sections, f"на участках: {on_sections}")
        self.assertIn("модель другого заказа («775_Другой»)", notices, "замечание: операции другого заказа — только в книге")
        loaded = self.s.sw.GetOpenDocumentByName(str(foreign))
        self.assertIsNotNone(loaded, "деталь другого заказа загружена сборкой")
        self.assertFalse(bool(com.dyn(loaded).GetSaveFlag), "деталь другого заказа не помечена изменённой")
        self.s.close_all()
        self.assertEqual(before, foreign.read_bytes(), "файл другого заказа не изменён")
        self.assertEqual(ops["ПРТИ.468211.111"], V(self.persisted(same), "Операции"), "деталь своего заказа — записана")

    def test_L06_assembly_in_any_folder(self):
        """L06 (замечание владельца 19.09.2026): главная сборка — в любой папке, вне заказа и без «01_3D»; детали — рядом и
        в другой папке. Книга ЛЗК — прямо рядом со сборкой, новых папок нет; повторный запуск заменяет книгу."""
        from eskd_e2e import build

        folder = self.case_dir / "Рабочая папка" / "Новый стол"
        folder.mkdir(parents=True, exist_ok=True)
        near = self.s.workspace_copy(Path(paths.FIXTURES_A) / SHEET_PART, subdir=f"{self._case_name()}/Рабочая папка/Новый стол")
        other = self.s.workspace_copy(Path(paths.FIXTURES_A) / "ПРТИ.468211.111 Стойка трубная.sldprt",
                                      subdir=f"{self._case_name()}/Детали где-то ещё")
        asm, opened = build.assembly(self.s, [(near, 0, 0, 0), (other, 0, 0.15, 0)])
        self.s.save_as(asm, folder / "Стол письменный.sldasm")
        for doc in opened:
            self.s.close(doc)
        self.s.activate(asm)
        status = self._build()
        notices = str(com.call(self.s.eskd(), "LastNotices"))
        self.path("outcome.txt").write_text(status + "\n\n" + notices, encoding="utf-8")
        self.assertTrue(status.startswith("ok|"), f"{status}\n{notices}")
        workbook = Path(status.split("|")[1])
        self.assertEqual(folder.resolve(), workbook.parent.resolve(), "книга прямо рядом со сборкой")
        self.assertEqual("ЛЗК_Стол письменный.xlsx", workbook.name)
        self.assertTrue(workbook.is_file(), "книга ЛЗК")
        ops, on_sections = self._sections(workbook)
        for designation in ("ПРТИ.468211.101", "ПРТИ.468211.111"):
            with self.subTest(part=designation):
                self.assertNotIn(ops.get(designation) or "?", ("?", ""), f"операции: {ops}\n{notices}")
                self.assertIn(designation, on_sections, f"на участках: {on_sections}")
        self.assertEqual([], [d.name for d in folder.iterdir() if d.is_dir()], "рядом со сборкой новых папок нет")
        status = self._build()
        self.assertTrue(status.startswith("ok|"), status)
        self.assertEqual([], [d.name for d in folder.iterdir() if d.is_dir()], "и после повторного запуска")
        self.assertEqual(["ЛЗК_Стол письменный.xlsx"], sorted(f.name for f in folder.glob("*.xlsx")))
        notices = str(com.call(self.s.eskd(), "LastNotices"))
        backup = next((line.split("её копия — ", 1)[1].strip() for line in notices.splitlines() if "её копия — " in line), "")
        self.assertTrue(backup and Path(backup).is_file(), f"копия прежней книги вне папки сборки\n{notices}")
        self.addCleanup(Path(backup).unlink, True)
        self.assertNotEqual(folder.resolve(), Path(backup).parent.resolve())

    def test_L08_unreadable_previous_book_is_not_overwritten(self):
        """L08: прежняя книга рядом со сборкой не читается — кнопка отказывает и её не трогает: введённое в ней
        (тираж, срок, нормы) не должно пропасть молча."""
        from eskd_e2e import build

        folder = self.case_dir / "Где угодно"
        folder.mkdir(parents=True, exist_ok=True)
        part = self.s.workspace_copy(Path(paths.FIXTURES_A) / SHEET_PART, subdir=f"{self._case_name()}/Где угодно")
        asm, _ = build.assembly(self.s, [(part, 0, 0, 0)])
        self.s.save_as(asm, folder / "Тумба.sldasm")
        broken = folder / "ЛЗК_Тумба.xlsx"
        broken.write_bytes(b"not a zip workbook")
        self.s.activate(asm)
        status = self._build()
        self.path("outcome.txt").write_text(status, encoding="utf-8")
        self.assertTrue(status.startswith("error|") and "не читается" in status, status)
        self.assertEqual(b"not a zip workbook", broken.read_bytes(), "прежняя книга не тронута")

    def test_L07_readonly_models_and_denied_folder_still_give_book(self):
        """L07: чужая сборка — деталь занята (только для чтения), в папку сборки писать нельзя. Раньше кнопка обрывалась
        «Ведомость не сформирована»; теперь книга — в «Документы», операции детали — в книге."""
        import subprocess
        from eskd_e2e import build

        stamp = time.strftime("%H%M%S")
        name = f"Стол чужой L07 {stamp}"
        folder = self.case_dir / "Чужая папка"
        folder.mkdir(parents=True, exist_ok=True)
        part = self.s.workspace_copy(Path(paths.FIXTURES_A) / SHEET_PART, subdir=f"{self._case_name()}/Чужая папка")
        asm, opened = build.assembly(self.s, [(part, 0, 0, 0)])
        asm_path = folder / f"{name}.sldasm"
        self.s.save_as(asm, asm_path)
        self.s.close_all()
        before = part.read_bytes()

        user = os.environ["USERNAME"]
        subprocess.run(["icacls", str(folder), "/deny", f"{user}:(AD,WD)"], check=True, capture_output=True)
        self.addCleanup(subprocess.run, ["icacls", str(folder), "/remove:d", user], capture_output=True)

        self.s.open(part, readonly=True)
        doc = self.s.open(asm_path)
        self.s.activate(doc)
        status = self._build()
        notices = str(com.call(self.s.eskd(), "LastNotices"))
        self.path("outcome.txt").write_text(status + "\n\n" + notices, encoding="utf-8")
        self.assertTrue(status.startswith("ok|"), f"{status}\n{notices}")
        workbook = Path(status.split("|")[1])
        self.assertEqual(f"ЛЗК_{name}.xlsx", workbook.name, "запасное место — «Документы», без папок")
        self.addCleanup(workbook.unlink, True)
        self.assertTrue(workbook.is_file(), "книга ЛЗК")
        self.assertIn("записать нельзя", notices)
        self.assertIn("только для чтения", notices)
        ops, on_sections = self._sections(workbook)
        self.assertNotIn(ops.get("ПРТИ.468211.101") or "?", ("?", ""), f"операции: {ops}\n{notices}")
        self.assertIn("ПРТИ.468211.101", on_sections)
        self.s.close_all()
        self.assertEqual(before, part.read_bytes(), "деталь только для чтения не изменена")

    def test_L05_base_parts_get_operations_only_in_book(self):
        """L05: детали из базы («_Библиотека проектирования») — операции по модели в книге и на участках, файл базы не меняется."""
        base = "_Библиотека проектирования/Общие детали"
        before = {}
        for n in (SHEET_PART, "ПРТИ.468211.111 Стойка трубная.sldprt"):
            src = Path(paths.FIXTURES_A) / n
            before[n] = src.read_bytes()
        product, parts, _, notices = self._outside(base)
        ops, on_sections = self._sections(product / DOCS / BOOK)
        for designation, path in (("ПРТИ.468211.101", parts[0]), ("ПРТИ.468211.111", parts[1])):
            with self.subTest(part=designation):
                self.assertNotIn(ops.get(designation) or "?", ("?", ""), f"операции детали базы: {ops}")
                self.assertIn(designation, on_sections, f"деталь базы на участках: {on_sections}")
                self.assertEqual(before[path.name], path.read_bytes(), "файл базы не изменён")
        self.assertIn("модель базы", notices, "пометка, что операции только в книге")

    @tags("smoke")
    def test_L01_builds_workbook_with_sheets(self):
        """L01: книга ЛЗК_<шифр>.xlsx в «04_Сопроводительная документация» — паспорт, участки, расход, нормы; «Операции» и «Габарит» в модели."""
        import openpyxl

        product, asm = self._product()
        doc = self.s.open(asm)
        self.s.activate(doc)
        status = self._build()
        self.assertTrue(status.startswith("ok|"), status)
        _, path, row_count, issues = status.split("|")
        workbook = product / DOCS / BOOK
        self.assertEqual(str(workbook).lower(), path.lower(), "имя и место ведомости")
        self.assertTrue(workbook.is_file(), "файл ведомости")
        self.assertFalse((product / "_Ведомость.txt").exists(), "отчёта-текстовика больше нет — замечания в окне")
        notices = str(com.call(self.s.eskd(), "LastNotices"))
        # Изделие здесь не проверялось: одно замечание — «книга собрана по непроверенному изделию» (З-27); пометки «?»
        # книги — замечания по строкам. Замечание о старой версии SWTools зависит от рабочего места и не считается.
        self.assertEqual(1, notices.count("непроверенному изделию"), f"книга без проверки — сказано: {notices}")
        self.assertEqual(int(issues), notices.count("ЗАМЕЧАНИЕ — Строка "), f"пометки «?» — замечания окна: {notices}")
        self.assertGreater(int(row_count), 5, "строк в ведомости")

        wb = openpyxl.load_workbook(workbook)
        self.assertEqual(SHEETS, wb.sheetnames)
        main = wb["Ведомость"]
        self.assertTrue(str(main["B2"].value or "").startswith("ПРТИ.468211.100"), main["B2"].value)
        self.assertEqual(ASM, main["B3"].value, "сборка в шапке")
        self.assertTrue(str(main["B4"].value or "").startswith("SHA-256 "), "контрольная сумма")
        rows = [r for r in range(7, main.max_row + 1) if main.cell(r, 10).value]
        self.assertEqual(int(row_count), len(rows), "строк в итоге кнопки и в книге")
        designations = [main.cell(r, 3).value for r in rows]
        self.assertIn("ПРТИ.468211.101", designations, "обозначение из свойства модели, без наименования")
        names = {main.cell(r, 3).value: main.cell(r, 4).value for r in rows}
        self.assertEqual("Пластина опорная", names["ПРТИ.468211.101"], "наименование")
        ops = {main.cell(r, 3).value: main.cell(r, 9).value for r in rows}
        self.assertEqual("Механическая сборка", ops["ПРТИ.468211.110"], "операции подсборки по признакам модели")
        self.assertTrue(all(ops.values()), f"у каждой строки операции или «?»: {ops}")
        # З-1: красится прокат; лист из стали — прокат, значит и резка, и покраска.
        self.assertIn("Покраска", ops["ПРТИ.468211.101"], f"стальной лист красится: {ops['ПРТИ.468211.101']}")
        self.assertNotIn("Покраска", ops["ПРТИ.468211.110"], "механическая сборка целиком не красится")
        paint = wb["Покрасочный"]
        self.assertEqual("Покрасочный участок", paint["A1"].value, "название листа")
        # В-3 (решение владельца 19.09.2026): подписей на листах участков нет
        for section in ("Заготовительный", "Сварочный", "Покрасочный", "Комплектовочный"):
            texts = [str(c.value) for row in wb[section].iter_rows() for c in row if isinstance(c.value, str)]
            self.assertFalse([t for t in texts if "_____ /" in t], f"{section}: строки подписей")
        self.assertEqual("№", paint["A5"].value, "шапка таблицы")
        self.assertEqual("=Тираж", paint["C3"].value, "тираж в шапке — из паспорта")
        data = [r for r in range(6, paint.max_row + 1) if isinstance(paint.cell(r, 1).value, (int, float))]
        painted = [paint.cell(r, 2).value for r in data]
        self.assertIn("ПРТИ.468211.101", painted, f"лист в покраске: {painted}")
        areas = [paint.cell(r, 4).value for r in data]
        self.assertTrue(all(isinstance(a, (int, float)) and a > 0 for a in areas), f"площадь каждой единицы: {areas}")
        self.assertEqual(len(set(areas)), len(areas), f"площади разных деталей различаются: {areas}")
        self.assertTrue(all(str(paint.cell(r, 6).value).endswith("*Тираж") for r in data), "всего на заказ — формулой")
        kit = wb["Комплектовочный"]
        codes = [kit.cell(r, 4).value for r in range(6, kit.max_row + 1) if isinstance(kit.cell(r, 1).value, (int, float))]
        self.assertNotIn("?", codes, f"нет кода — ячейка пустая, а не «?»: {codes}")
        self.assertTrue(str(main.auto_filter.ref or "").startswith("A6:M"), f"фильтр по таблице с категорией: {main.auto_filter.ref}")
        self.assertTrue(all(main.cell(r, 11).value for r in rows), "категория у каждой строки")
        self.assertIn("Тираж", wb.defined_names, "имя «Тираж»")
        self.assertIn("Хлыст", wb.defined_names, "норматив «Хлыст»")
        self.assertFalse((product / "Ведомость_ПРТИ.468211.100.xlsx").exists(), "в корне изделия ведомости нет")
        sizes = [main.cell(r, 6).value for r in rows]
        self.assertTrue(all(sizes), f"габарит у каждой строки: {sizes}")
        self.assertEqual(int(issues), int(str(main["G4"].value).split()[0]) if main["G4"].value != "нет" else 0,
                         "число замечаний в шапке")

        self.s.close_all()
        self.assertEqual("Механическая сборка", V(self.persisted(asm.parent / SUBASSEMBLY), "Операции"), "«Операции» в подсборке")
        self.assertEqual("200×100×4", V(self.persisted(asm.parent / SHEET_PART), "Габарит"), "«Габарит» в детали")
        self.assertEqual([], self.addin_errors(), "ошибки в журнале надстройки")

    def test_L02_second_run_archives_previous(self):
        """L02: повторное формирование переносит прежнюю книгу в _Аннулировано, введённый тираж и нормы сохраняет;
        ведомость старого образца из корня изделия тоже уходит в _Аннулировано."""
        import openpyxl

        product, asm = self._product()
        legacy = product / "Ведомость_ПРТИ.468211.100.xlsx"
        openpyxl.Workbook().save(legacy)
        doc = self.s.open(asm)
        self.s.activate(doc)
        self.assertTrue(self._build().startswith("ok|"))
        self.assertFalse(legacy.exists(), "ведомость старого образца убрана из корня")
        self.assertEqual(1, len(list((product / "_Аннулировано").glob("Ведомость_ПРТИ.468211.100_*.xlsx"))), "старый образец в архиве")

        # Начальник производства ввёл тираж и правку нормы — повторная К-4 их не теряет.
        book = product / DOCS / BOOK
        wb = openpyxl.load_workbook(book)
        sheet, cell = next(iter(wb.defined_names["Тираж"].destinations))
        wb[sheet][cell.replace("$", "")] = 41
        sheet, cell = next(iter(wb.defined_names["Захват"].destinations))
        wb[sheet][cell.replace("$", "")] = 600
        wb.save(book)
        self.assertTrue(self._build().startswith("ok|"))
        archived = list((product / "_Аннулировано").glob("ЛЗК_ПРТИ.468211.100_*.xlsx"))
        self.assertEqual(1, len(archived), archived)
        wb = openpyxl.load_workbook(book)
        sheet, cell = next(iter(wb.defined_names["Тираж"].destinations))
        self.assertEqual(41, wb[sheet][cell.replace("$", "")].value, "тираж сохранён")
        sheet, cell = next(iter(wb.defined_names["Захват"].destinations))
        self.assertEqual(600, wb[sheet][cell.replace("$", "")].value, "норма сохранена")

    def test_L09_bar_without_cut_list_length_is_measured_along_its_body(self):
        """L09 (решение владельца 21.09.2026): труба 500 мм под 53° к осям детали, длины в списке вырезов нет —
        так SolidWorks ведёт себя, когда тело элемента конструкции дорабатывали. Габарит по осям дал бы 424 мм,
        и в закуп ушёл бы короткий хлыст. Книга меряет тело вдоль его самого длинного прямого ребра — 500 —
        и честно помечает длину «*»."""
        import openpyxl
        from eskd_e2e import build

        product = self.case_dir / PRODUCT
        (product / "01_3D").mkdir(parents=True, exist_ok=True)
        for name in ("Нормативы_производства.xlsx", "ЛЗК_бланки.xlsx"):
            ref = paths.ROOT / "02_Шаблоны_и_Форматки" / "Справочники" / name
            (self.case_dir / ref.name).write_bytes(ref.read_bytes())
        tube = product / "01_3D" / "ПРТИ.468211.171 Раскос.sldprt"
        doc = build.structural_tube(self.s, 500, build.tube_profile(),
                                    "Труба 40х40х2,0 ГОСТ 8639-82 / Ст3сп ГОСТ 13663-86", angle_deg=53.13)
        self.s.save_as(doc, tube)
        self.s.wait_addin_idle(timeout=60.0)
        box = com.call(doc, "GetPartBox", True)
        self.assertLess(max((box[3] - box[0]), (box[4] - box[1]), (box[5] - box[2])) * 1000, 450,
                        "наклонная труба: ни один габарит по осям не дотягивает до 500")
        asm, opened = build.assembly(self.s, [(tube, 0, 0, 0)])
        self.s.save_as(asm, product / "01_3D" / ASM)
        for d in opened:
            self.s.close(d)
        self.s.activate(asm)
        status = self._build()
        notices = str(com.call(self.s.eskd(), "LastNotices"))
        self.assertTrue(status.startswith("ok|"), f"{status}\n{notices}")

        main = openpyxl.load_workbook(product / DOCS / BOOK)["Ведомость"]
        sizes = {main.cell(r, 3).value: str(main.cell(r, 6).value or "") for r in range(7, main.max_row + 1)}
        size = sizes.get("ПРТИ.468211.171")
        self.assertIsNotNone(size, f"раскос в ведомости: {sizes}")
        self.assertRegex(size, r"^L=\d+(?:[.,]\d+)?\*$", f"длина, а не габарит «Д×Ш×В», и со звёздочкой: {size}")
        self.assertEqual("L=500*", size, "вдоль собственной оси — 500, а не 424 по осям детали")
        self.assertIn("измерена по модели", notices, f"пометка объясняет, откуда длина: {notices}")

    def test_L10_bent_sheet_part_gets_flat_pattern_size(self):
        """L10 (замечание владельца 21.09.2026): у гнутой листовой детали заготовка — развёртка, а не габарит готовой
        детали. Уголок 100×25 с отгибом 20 из листа 3 мм (радиус гиба 3): согнутый — 103×25×23, развёртка —
        121,07×25×3 (полки 97 и 17 плюс дуга по нейтральному слою 7,07)."""
        import openpyxl
        from eskd_e2e import build

        product = self.case_dir / PRODUCT
        (product / "01_3D").mkdir(parents=True, exist_ok=True)
        for name in ("Нормативы_производства.xlsx", "ЛЗК_бланки.xlsx"):
            ref = paths.ROOT / "02_Шаблоны_и_Форматки" / "Справочники" / name
            (self.case_dir / ref.name).write_bytes(ref.read_bytes())
        angle = product / "01_3D" / "ПРТИ.468211.172 Закладная.sldprt"
        doc = build.sheet_metal_angle(self.s, 100, 20, 25, 3, "Лист 3,0 ГОСТ 19903-2015 / Ст3сп ГОСТ 14637-89")
        self.s.save_as(doc, angle)
        self.s.wait_addin_idle(timeout=60.0)
        box = com.call(doc, "GetPartBox", True)
        folded = sorted(((box[i + 3] - box[i]) * 1000 for i in range(3)), reverse=True)
        self.assertAlmostEqual(103, folded[0], delta=0.5, msg=f"согнутая деталь по осям: {folded}")
        self.assertAlmostEqual(25, folded[1], delta=0.5, msg=f"согнутая деталь по осям: {folded}")
        self.assertAlmostEqual(23, folded[2], delta=0.5, msg=f"согнутая деталь по осям: {folded}")
        asm, opened = build.assembly(self.s, [(angle, 0, 0, 0)])
        self.s.save_as(asm, product / "01_3D" / ASM)
        for d in opened:
            self.s.close(d)
        self.s.activate(asm)
        status = self._build()
        notices = str(com.call(self.s.eskd(), "LastNotices"))
        self.assertTrue(status.startswith("ok|"), f"{status} {notices}")

        main = openpyxl.load_workbook(product / DOCS / BOOK)["Ведомость"]
        sizes = {main.cell(r, 3).value: str(main.cell(r, 6).value or "") for r in range(7, main.max_row + 1)}
        size = sizes.get("ПРТИ.468211.172")
        self.assertIsNotNone(size, f"закладная в ведомости: {sizes}")
        self.assertRegex(size, r"^\d+(?:[.,]\d+)?×\d+(?:[.,]\d+)?×\d+(?:[.,]\d+)?$", f"развёртка Д×Ш×S без «*»: {size}")
        length, width, thickness = (float(x.replace(",", ".")) for x in size.split("×"))
        self.assertAlmostEqual(121.07, length, delta=0.1, msg=f"длина развёртки — полки и дуга, а не габарит 103: {size}")
        self.assertAlmostEqual(25, width, delta=0.5, msg=f"ширина развёртки — глубина уголка: {size}")
        self.assertAlmostEqual(3, thickness, delta=0.05, msg=f"третий размер — толщина листа: {size}")
        self.s.close_all()
        self.assertEqual(size, V(self.persisted(angle), "Габарит"), "«Габарит» в детали — тоже развёртка")

    def test_L11_variants_get_own_size_and_top_is_not_kitted(self):
        """L11 (аудит и замечание владельца 21.09.2026): у двух исполнений одного файла разная толщина — в книге у каждого
        свой размер, а не размер активной конфигурации. Главная сборка с «Механической сборкой» не попадает на
        «Комплектовочный»: комплектуют то, из чего её собирают, а не её саму."""
        import openpyxl
        from eskd_e2e import build

        product = self.case_dir / PRODUCT
        (product / "01_3D").mkdir(parents=True, exist_ok=True)
        for name in ("Нормативы_производства.xlsx", "ЛЗК_бланки.xlsx"):
            ref = paths.ROOT / "02_Шаблоны_и_Форматки" / "Справочники" / name
            (self.case_dir / ref.name).write_bytes(ref.read_bytes())
        plate = product / "01_3D" / "ПРТИ.468211.173 Прокладка.sldprt"
        doc, feat = build.plate(self.s, 150, 80, 4, "Лист 4,0 ГОСТ 19903-2015 / Ст3сп ГОСТ 14637-89")
        active = str(doc.GetActiveConfiguration.Name)
        build.add_configuration(doc, "01")
        # толщина меняется только в «01»: вариант 1 — «только в этой», то есть в активной. Список конфигураций
        # (вариант 3) через позднее связывание COM не доходит — GetSystemValue3(3, [...]) возвращает None
        self.s.activate(doc)
        build.show_configuration(doc, "01")
        self.assertEqual("01", str(doc.GetActiveConfiguration.Name), "размер задаётся в «01»")
        rc = com.dyn(doc.Parameter("D1@" + str(feat.Name))).SetSystemValue3(0.010, 1, None)
        self.assertEqual(0, rc, "толщина исполнения 01 задана")
        doc.ForceRebuild3(False)
        box = com.call(doc, "GetPartBox", True)
        dim = com.dyn(doc.Parameter("D1@" + str(feat.Name)))
        diag = (f"feat={feat.Name} active={doc.GetActiveConfiguration.Name} "
                f"cfgs={list(com.as_list(doc.GetConfigurationNames))} "
                f"vals={com.as_list(dim.GetSystemValue3(2, None))} box={list(box)}")
        self.assertAlmostEqual(10, min(abs(box[i + 3] - box[i]) for i in range(3)) * 1000, delta=0.1,
                               msg=f"в «01» толщина 10: {diag}")
        build.show_configuration(doc, active)
        doc.ForceRebuild3(False)
        self.s.save_as(doc, plate)
        self.s.wait_addin_idle(timeout=60.0)
        asm, opened = build.assembly(self.s, [(plate, 0, 0, 0), (plate, 0, 0.2, 0)])
        comps = com.as_list(asm.GetComponents(True))
        com.dyn(comps[1]).ReferencedConfiguration = "01"
        asm.ForceRebuild3(False)
        build.props(asm, {"Операции": "Механическая сборка"}, "")
        self.s.save_as(asm, product / "01_3D" / ASM)
        for d in opened:
            self.s.close(d)
        self.s.activate(asm)
        status = self._build()
        notices = str(com.call(self.s.eskd(), "LastNotices"))
        self.assertTrue(status.startswith("ok|"), f"{status}\n{notices}")

        wb = openpyxl.load_workbook(product / DOCS / BOOK)
        main = wb["Ведомость"]
        sizes = {str(main.cell(r, 3).value): str(main.cell(r, 6).value or "") for r in range(7, main.max_row + 1)}
        self.assertEqual("150×80×4", sizes.get("ПРТИ.468211.173"), f"исполнение 00: {sizes}")
        self.assertEqual("150×80×10", sizes.get("ПРТИ.468211.173-01"), f"исполнение 01 — своя толщина: {sizes}")
        kit = wb["Комплектовочный"]
        kitted = [str(kit.cell(r, 2).value or "") for r in range(6, kit.max_row + 1)]
        self.assertFalse(any(v.startswith("ПРТИ.468211.100") for v in kitted), f"главная сборка не комплектуется: {kitted}")
        self.s.close_all()

    # ------------------------------------------------------------------ версия изделия (З-27)
    def _checked(self):
        """Проверка изделия без окна; возвращает (итог, текст отчёта, версия)."""
        import re

        com.call(self.s.eskd(), "CheckProductSilent")
        status = str(com.call(self.s.eskd(), "CheckStatus"))
        self.assertTrue(status.startswith("ok|"), status)
        text = Path(status.split("|")[4]).read_text(encoding="utf-8-sig")
        found = re.search(r"^Версия:\s+(.+)$", text, re.M)
        self.assertIsNotNone(found, f"в отчёте проверки есть версия изделия:\n{text}")
        return status, text, found.group(1).strip()

    @staticmethod
    def _book_version(workbook):
        import openpyxl
        wb = openpyxl.load_workbook(workbook)
        ws, cell = next(iter(wb.defined_names["Паспорт_Версия"].destinations))
        return str(wb[ws][cell.replace("$", "")].value or "").strip()

    def test_L13_version_lzk_export(self):
        """L13 (решение владельца 23.09.2026, З-27): проверка пишет версию изделия; ЛЗК и выгрузка после неё ничего не
        спрашивают и запоминают эту версию; собственные сохранения ЛЗК (записанные «Операции») версию не меняют —
        повторная проверка оставляет прежнюю, и замечаний о версии книги и выгрузки нет (сохранения выгрузки — X09).
        Несохранённая правка версию не меняет, а даёт замечание проверки. Деталь изменили и сохранили — новая версия,
        книга и выгрузка «по другой версии изделия». Имя теста короткое: у фикстуры есть файл с именем в 82 знака, и с
        длинным каталогом путь переваливает за 260."""
        product, asm = self._product()
        doc = self.s.open(asm)
        self.s.activate(doc)
        _, _, version = self._checked()

        lzk = self._build()
        self.assertTrue(lzk.startswith("ok|"), lzk)
        notices = str(com.call(self.s.eskd(), "LastNotices"))
        self.assertNotIn("непроверенному изделию", notices, f"изделие проверено и не менялось: {notices}")
        self.assertEqual(version, self._book_version(product / DOCS / BOOK), "книга помнит версию проверки")
        report = (product / "_Проверка.txt").read_text(encoding="utf-8-sig")
        self.assertIn("Версия:   " + version, report, "версия в отчёте проверки прежняя")
        self.assertIn("Суммы обновлены: ведомость ЛЗК", report, "ЛЗК вписала суммы моделей, которые сохранила сама")
        self.assertFalse(bool(doc.GetSaveFlag), "сборка после ЛЗК не помечена изменённой")

        com.call(self.s.eskd(), "ExportProductSilent")
        export = str(com.call(self.s.eskd(), "ExportStatus"))
        self.assertTrue(export.startswith("ok|"), export)
        exported = Path(export.split("|")[3]).read_text(encoding="utf-8-sig")
        self.path("export.txt").write_text(exported, encoding="utf-8")
        self.assertIn("Версия:   " + version, exported, "выгрузка помнит версию проверки")
        self.assertNotIn("непроверенному изделию", exported, exported)

        status, text, again = self._checked()
        self.path("check2.txt").write_text(text, encoding="utf-8")
        self.assertEqual(version, again, "свои сохранения ЛЗК и выгрузки версию не меняют")
        self.assertNotIn("другой версии", text, text)
        self.assertNotIn("непроверенному", text, text)
        self.assertNotIn("без отметки версии", text, text)

        # Деталь уже загружена сборкой: берётся открытая модель компонента.
        part = com.dyn(self.s.sw.GetOpenDocumentByName(str(asm.parent / SHEET_PART)))
        from eskd_e2e import build
        build.props(part, {"Примечание": "изменено после проверки"})
        part.SetSaveFlag()
        # Несохранённая правка: файлы прежние — версия прежняя (правку могут закрыть без сохранения), но проверка говорит
        # о ней, и «Готово к производству» не пройдёт, пока её не сохранят.
        _, text, same = self._checked()
        self.path("check3.txt").write_text(text, encoding="utf-8")
        self.assertEqual(version, same, "несохранённая правка версию не меняет")
        self.assertIn(SHEET_PART.lower().rsplit(".", 1)[0], text.lower(), text)
        self.assertIn("несохранённые правки", text, text)
        self.assertNotIn("другой версии", text, text)

        self.assertTrue(self.s.save(part)[0], "деталь сохранена")
        self.s.wait_addin_idle(timeout=60.0)
        _, text, changed = self._checked()
        self.path("check4.txt").write_text(text, encoding="utf-8")
        self.assertNotEqual(version, changed, "деталь изменили — новая версия")
        self.assertIn("книга собрана по другой версии изделия", text, text)
        self.assertIn("выгрузка сделана по другой версии изделия", text, text)
        self.assertEqual([], self.addin_errors(), "ошибки в журнале надстройки")

    def test_L14_edited_models_not_saved(self):
        """L14 (решение владельца 23.09.2026, З-25, З-27): у детали несохранённые правки конструктора — ЛЗК пишет в неё
        «Операции», но не сохраняет; изделие не проверено — книга собирается «как есть» с отметкой «не проверено»."""
        product, asm = self._product()
        doc = self.s.open(asm)
        part = com.dyn(self.s.sw.GetOpenDocumentByName(str(asm.parent / SHEET_PART)))
        from eskd_e2e import build
        build.props(part, {"Примечание": "правка конструктора"})
        part.SetSaveFlag()
        self.assertTrue(bool(part.GetSaveFlag), "у детали несохранённые правки")
        before = (asm.parent / SHEET_PART).read_bytes()
        self.s.activate(doc)
        status = self._build()
        notices = str(com.call(self.s.eskd(), "LastNotices"))
        self.path("outcome.txt").write_text(status + "\n\n" + notices, encoding="utf-8")
        self.assertTrue(status.startswith("ok|"), f"{status}\n{notices}")
        self.assertEqual(before, (asm.parent / SHEET_PART).read_bytes(), "файл детали с правками конструктора не сохранён")
        self.assertTrue(bool(part.GetSaveFlag), "правки конструктора по-прежнему несохранённые")
        self.assertIn("несохранённые правки", notices, notices)
        self.assertIn("непроверенному изделию", notices, "изделие не проверялось — сказано")
        self.assertEqual("не проверено", self._book_version(product / DOCS / BOOK), "паспорт: книга по непроверенному изделию")
        self.assertEqual([], self.addin_errors(), "ошибки в журнале надстройки")

    def test_L03_refuses_non_assembly(self):
        """L03: у детали кнопка отказывает без изменений файлов."""
        path, doc = self.open_copy(SHEET_PART)
        self.s.activate(doc)
        status = self._build()
        self.assertTrue(status.startswith("error|"), status)
        self.assertNoPropertyWrites()


if __name__ == "__main__":
    unittest.main()
