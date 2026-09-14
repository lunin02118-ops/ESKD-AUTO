# -*- coding: utf-8 -*-
"""Спайк С-6: «шаблон заказчика» — стили, объединения, ширины, заголовки не в первой строке, второй лист, формула."""
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

out = sys.argv[1]
wb = Workbook()
ws = wb.active
ws.title = "СЗ"
ws["A1"] = "Служебная записка № 52 от 15.09.2026"
ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:J1")
ws["A2"] = "Заказчик: ТОО «Тест»"
headers = [" позиция", "Изделие", "СТАНДАРТНОСТЬ", "Направление", "Ш и ф р", "Эталон", "Исполнение", "Тираж", "RAL", "Поверхность"]
thin = Side(style="thin")
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=3, column=i, value=h)
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
    c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.column_dimensions[c.column_letter].width = 16
rows = [
    (1, "Стеллаж", "СТД", "МЕТАЛЛ", "ТС-52-С1", "ТС-00-С1", "", 10, "", ""),
    (2, "Стол", "ТИП", "КОМБИ", "ТС-52-Т1", "ТС-00-Т1", "-01", 2, "RAL 9005", "шагрень"),
    (3, "Тумба", "НОВ", "КОРПУС", "ТС-52-Н1", "", "", 1, "", ""),
]
for r, row in enumerate(rows, start=4):
    for i, v in enumerate(row, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
ws["I5"].fill = PatternFill("solid", fgColor="FFF2CC")
ws["K4"] = "=H4*2"
wb.create_sheet("Справка")["A1"] = "второй лист не трогается"
wb.save(out)
print("шаблон:", out)
