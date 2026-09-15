# -*- coding: utf-8 -*-
"""Макеты Excel для кнопок ТЗ-02 (этап Э0, Т-11, Т-12, Т-51) до выдачи шаблонов заказчиком.

Кнопки читают эти файлы по заголовкам столбцов (Т-10), поэтому шаблон заказчика может заменить макет без правки кода,
если в нём есть те же заголовки. Разделы СЗ на производство — отдельные листы: так их проще читать и дописывать.

    python 09_Тесты/tools/make_tz02_templates.py          — записать макеты в 02_Шаблоны_и_Форматки/Шаблоны документов
    python 09_Тесты/tools/make_tz02_templates.py --prod заявка.json заявка.xlsx  — заполненная СЗ на производство
"""
import sys
from pathlib import Path

import json

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from tz02_norms import norms_workbook  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "02_Шаблоны_и_Форматки" / "Шаблоны документов"
NORMS_PATH = ROOT / "02_Шаблоны_и_Форматки" / "Справочники" / "Нормативы_производства.xlsx"
THIN = Side(style="thin", color="808080")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
FONT = "Arial"


def header_row(ws, row, columns, note_columns=()):
    for i, (title, width) in enumerate(columns, start=1):
        c = ws.cell(row=row, column=i, value=title)
        c.font = Font(name=FONT, bold=True, size=10)
        c.fill = NOTE_FILL if title in note_columns else HEAD_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def body_rows(ws, first, count, width):
    for r in range(first, first + count):
        for col in range(1, width + 1):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            c.font = Font(name=FONT, size=10)


def title(ws, text, fields):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    for i, (label, value) in enumerate(fields, start=2):
        ws.cell(row=i, column=1, value=label).font = Font(name=FONT, bold=True, size=10)
        ws.cell(row=i, column=2, value=value).font = Font(name=FONT, size=10)


def order_note():
    wb = Workbook()
    ws = wb.active
    ws.title = "СЗ"
    title(ws, "Служебная записка на изготовление", [("№ заказа:", ""), ("Заказчик:", ""), ("Дата:", ""), ("Срок:", "")])
    columns = [("Позиция", 9), ("Изделие", 32), ("Стандартность", 14), ("Направление", 13), ("Шифр", 18), ("Эталон", 18),
               ("Исполнение", 12), ("Тираж", 9), ("RAL", 14), ("Поверхность", 14), ("Примечание", 30)]
    header_row(ws, 7, columns)
    body_rows(ws, 8, 30, len(columns))
    for col, values in ((3, "СТД,ТИП,НОВ"), (4, "МЕТАЛЛ,КОРПУС,КОМБИ"), (10, "глянец,матовый,шагрень,муар")):
        dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
        letter = get_column_letter(col)
        dv.add(f"{letter}8:{letter}200")
        ws.add_data_validation(dv)
    dv = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True)
    dv.add("H8:H200")
    ws.add_data_validation(dv)
    help_ws = wb.create_sheet("Справка")
    rows = [
        ("Столбец", "Что писать", "Кто читает"),
        ("Позиция", "номер строки служебной записки", "К-5 «Выдать в производство»"),
        ("Изделие", "наименование, как у заказчика", "человек"),
        ("Стандартность", "СТД — эталон из 02_БАЗА без изменений; ТИП — эталон с небольшими изменениями; НОВ — новое", "К-5: где искать папку изделия"),
        ("Направление", "МЕТАЛЛ, КОРПУС или КОМБИ (металл + корпус)", "К-5: папка направления"),
        ("Шифр", "шифр изделия в заказе; имя папки изделия", "К-5: сопоставление с папкой"),
        ("Эталон", "шифр эталона в 02_БАЗА (для СТД и ТИП)", "К-8: применяемость эталона"),
        ("Исполнение", "исполнение эталона, например -01", "К-5"),
        ("Тираж", "количество изделий, целое", "К-5: СЗ на производство × тираж"),
        ("RAL", "цвет покрытия, например RAL 7024; пусто — без покрытия", "К-4, К-5: раздел «Покрасить»"),
        ("Поверхность", "глянец, матовый, шагрень, муар", "К-4, К-5"),
        ("Примечание", "свободный текст", "человек"),
        ("", "Столбцы находятся по заголовкам в первых 10 строках: порядок и регистр не важны. Лишние столбцы можно "
             "добавлять. Кнопка может дописать скрытый столбец «Папка» — запомненный ответ сопоставления.", ""),
    ]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            cell = help_ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, bold=(r == 1), size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for letter, width in (("A", 16), ("B", 70), ("C", 34)):
        help_ws.column_dimensions[letter].width = width
    return wb


# СЗ на производство повторяет сводную заявку КТО (85_Т, 04.09.2026): пять листов по участкам, печать на A4 книжной,
# разметка и содержимое закрыты паролем. Реквизиты — на скрытом листе «Данные», шапки всех листов ссылаются на него.
PROD_PASSWORD = "kto"
BLUE = "1F4FD8"
GREY_FILL = PatternFill("solid", fgColor="E3E7ED")
QTY = "Данные!$B$7"
DATA_FIELDS = [("Фирма", "ТОО «Компания TROYA»"), ("Дата", ""), ("№ заявки", ""), ("Нач. КТО", ""), ("Изделие", ""),
               ("Шифр", ""), ("Кол-во", None)]
# Формат чисел: ноль не печатается, чтобы пустые строки бланка оставались пустыми.
INT, DEC1, DEC2 = "0;-0;;@", "0.0;-0.0;;@", "0.00;-0.00;;@"

# Столбец: (заголовок, ширина, формат, выравнивание, формула «Всего»). В формуле {r} — строка, {q} — тираж.
PROD_SECTIONS = [
    {
        "sheet": "1 Заготовка",
        "title": "1. ЦЕХ МЕТАЛЛОИЗДЕЛИЙ (Заготовительный участок: лазерная резка, гибка)",
        "subtitle": "Спецификация чистых заготовок на тираж (лазерная резка листа и трубы, гибка)",
        "columns": [("№", 4, INT, "center", None), ("Обозначение", 16, None, "center", None),
                    ("Наименование детали", 19, None, "left", None), ("Сортамент / Материал", 25, None, "left", None),
                    ("L, мм", 9, DEC1, "right", None), ("На 1", 6, INT, "right", None),
                    ("Всего, шт", 8, INT, "right", '=IF(F{r}="","",F{r}*{q})'), ("Операция", 22, None, "left", None)],
        "total": {"label": '="ИТОГО ЗАГОТОВОК НА ЗАКАЗ ("&' + QTY + '&" шт.)"', "span": 5,
                  "cells": {6: "=SUM({c})", 7: "=SUM({c})", 8: '=IF(COUNTA(B{a}:B{b})=0,"",COUNTA(B{a}:B{b})&" типоразм.")'}},
        "notes": ["Трубы и профили резать на лазерном труборезе по файлам IGS, листовые детали — на лазере по чистому контуру DXF из папки 03_ЧПУ изделия.",
                  "Перед передачей в сварочный цех снять заусенцы на торцах всех заготовок и зачистить зоны сварных стыков."],
        "signers": ["Начальник цеха металлоизделий"],
    },
    {
        "sheet": "2 Сварка",
        "title": "2. СВАРОЧНО-СБОРОЧНЫЙ ЦЕХ",
        "subtitle": "Сборочные единицы: сварная и механическая сборка, масса конструкций",
        "columns": [("№", 4, INT, "center", None), ("Обозначение узла", 16, None, "center", None),
                    ("Наименование сборочной единицы", 24, None, "left", None), ("Операция", 17, None, "left", None),
                    ("Масса 1 шт, кг", 10, DEC2, "right", None),
                    ("На 1 изд", 8, INT, "right", None), ("Всего на тираж", 10, INT, "right", '=IF(F{r}="","",F{r}*{q})'),
                    ("Общая масса, кг", 12, DEC2, "right", '=IF(G{r}="","",E{r}*G{r})')],
        "total": {"label": "ИТОГО СБОРОЧНЫХ ЕДИНИЦ НА ЗАКАЗ", "span": 5,
                  "cells": {6: "=SUM({c})", 7: "=SUM({c})", 8: "=SUM({c})"}},
        "notes": ["Сварка полуавтоматическая в среде защитных газов (80% Ar + 20% CO2) сварочной проволокой Св-08Г2С ø0.8 мм по ГОСТ 14771-76.",
                  "Сварные швы сплошные, катет шва 2.0...3.0 мм. Не допускаются непровары, прожоги тонкостенных труб и наплывы металла.",
                  "Все лицевые сварные швы зачистить лепестковым кругом заподлицо с плоскостью труб под полимерную покраску."],
        "signers": ["Начальник сварочного цеха"],
    },
    {
        "sheet": "3 Покраска",
        "title": "3. УЧАСТОК ПОЛИМЕРНО-ПОРОШКОВОЙ ПОКРАСКИ",
        "subtitle": "Ведомость окрашиваемых сборочных единиц и цвет полимерного покрытия",
        "columns": [("№", 4, INT, "center", None), ("Обозначение узла", 17, None, "center", None),
                    ("Наименование сборочной единицы", 36, None, "left", None), ("На 1 изд", 9, INT, "right", None),
                    ("Всего на тираж", 11, INT, "right", '=IF(D{r}="","",D{r}*{q})'), ("Цвет покрытия", 19, None, "center", None)],
        "total": {"label": "ИТОГО НА ПОКРАСКУ", "span": 3, "cells": {4: "=SUM({c})", 5: "=SUM({c})"}},
        "notes": ["Требования к внешнему виду: покрытие должно быть сплошным, однородным, без непрокрасов, потеков, царапин и механических повреждений.",
                  "Контроль качества внешнего вида покрытия производить визуально при естественном или рассеянном искусственном освещении.",
                  "Окрашенные сборочные единицы остудить и передать на участок комплектации, фурнитуры и сборки."],
        "signers": ["Мастер малярного участка"],
    },
    {
        "sheet": "4 Комплектация",
        "title": "4. СКЛАД ТМЦ И УЧАСТОК КОМПЛЕКТАЦИИ (Фурнитура и покупные изделия)",
        "subtitle": "Комплектовочная ведомость покупной фурнитуры, комплектующих и настила",
        "columns": [("№", 4, INT, "center", None), ("Код 1С", 13, None, "center", None),
                    ("Номенклатура ТМЦ", 32, None, "left", None), ("Ед. изм.", 7, None, "center", None),
                    ("На 1 изд", 8, "General", "right", None),
                    ("Всего на тираж", 10, "General", "right", '=IF(E{r}="","",E{r}*{q})'),
                    ("Назначение", 24, None, "left", None)],
        "total": None,
        "notes": ["Покупные изделия выдаются со склада ТМЦ на участки согласно графе «Назначение»."],
        "signers": ["Заведующий складом ТМЦ"],
    },
    {
        "sheet": "5 Списание",
        "title": "5. СВОДНАЯ ВЕДОМОСТЬ РАСХОДА И СПИСАНИЯ ТМЦ",
        "subtitle": "Расчет потребности в хлыстах, КИМ, деловых обрезков и технологического отхода по справочнику «Нормативы производства»",
        "columns": [("№", 4, INT, "center", None), ("Сортамент и ГОСТ", 25, None, "left", None),
                    ("Чистый расход, м", 9, DEC1, "right", None), ("Хлыстов", 8, None, "right", None),
                    ("Расход с уч. техн., м", 9, DEC1, "right", None), ("Расход с уч. техн., кг", 9, DEC1, "right", None),
                    ("КИМ, %", 7, "0.0%;-0.0%;;@", "right", None), ("Деловые обрезки", 25, None, "left", None),
                    ("Отход, кг", 8, DEC1, "right", None)],
        "total": {"label": "ИТОГО МЕТАЛЛОПРОКАТ НА ЗАКАЗ", "span": 2,
                  "cells": {3: "=SUM({c})", 5: "=SUM({c})", 6: "=SUM({c})", 9: "=SUM({c})"}},
        "notes": ["Технологический отход металла (торцовка, рез и обрезки короче делового) подлежит сдаче в металлолом."],
        "signers": ["Начальник производства", "Заведующий складом ТМЦ"],
    },
]


def text_height(text, width_chars, size=10, factor=1.0):
    """Высота строки под перенос: Excel не подбирает высоту объединённых ячеек, а на защищённом листе её не поправить."""
    import math
    lines = sum(max(1, math.ceil(len(part) * factor / max(width_chars, 1))) for part in str(text).split("\n"))
    return round(lines * size * 1.32 + 3, 1)


def prod_sheet(wb, spec, data, section_data):
    ws = wb.create_sheet(spec["sheet"])
    cols = spec["columns"]
    n = len(cols)
    last = get_column_letter(n)
    total_width = sum(c[1] for c in cols)
    for i, col in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = col[1]
    ws.sheet_view.showGridLines = False
    bold = Font(name=FONT, bold=True, size=11)

    # Слева в первой строке длинное название, справа в третьей — шифр с тиражом: граница блоков своя у каждой строки.
    head = [('=Данные!B1&" | Производственная заявка"', '="Дата: "&Данные!B2', 22),
            ('="№ заявки: "&Данные!B3', '="Нач. КТО: "&Данные!B4', 30),
            ('="Изделие: "&Данные!B5', '="Шифр: "&Данные!B6&"  |  Кол-во: "&Данные!B7&" шт."', 42)]
    for r, (left, right, right_width) in enumerate(head, start=1):
        split = n - 1
        while split > 1 and sum(c[1] for c in cols[split:]) < right_width:
            split -= 1
        left_end = get_column_letter(split)
        right_start = get_column_letter(split + 1)
        ws.merge_cells(f"A{r}:{left_end}{r}")
        ws.merge_cells(f"{right_start}{r}:{last}{r}")
        ws[f"A{r}"] = left
        ws[f"A{r}"].font = Font(name=FONT, bold=True, size=13 if r == 1 else 11)
        ws[f"{right_start}{r}"] = right
        ws[f"{right_start}{r}"].font = bold
        ws[f"{right_start}{r}"].alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[r].height = 20
    ws.row_dimensions[4].height = 6

    number = PROD_SECTIONS.index(spec) + 1
    ws.merge_cells(f"A5:{last}5")
    ws["A5"] = CellRichText(TextBlock(InlineFont(rFont=FONT, sz=12, b=True, color=BLUE), f"ЛИСТ {number}  "),
                            TextBlock(InlineFont(rFont=FONT, sz=12, b=True), spec["title"]))
    ws.row_dimensions[5].height = text_height("ЛИСТ 1  " + spec["title"], total_width * 0.85, 12)
    ws.merge_cells(f"A6:{last}6")
    ws["A6"] = spec["subtitle"]
    ws["A6"].font = Font(name=FONT, italic=True, size=9, color="404040")

    for i, col in enumerate(cols, start=1):
        c = ws.cell(row=7, column=i, value=col[0])
        c.font = Font(name=FONT, bold=True, size=9)
        c.fill = GREY_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[7].height = max(text_height(col[0], col[1], 9) for col in cols)

    rows = section_data.get("rows") or [[None] * n for _ in range(5)]
    first = 8
    for k, values in enumerate(rows):
        r = first + k
        height = 14
        for i, col in enumerate(cols, start=1):
            value = values[i - 1] if i - 1 < len(values) else None
            if value is None and col[4]:
                value = col[4].format(r=r, q=QTY)
            if i == 1 and section_data.get("rows"):
                value = k + 1
            c = ws.cell(row=r, column=i, value=value)
            c.font = Font(name=FONT, size=9, bold=(col[4] is not None and i == n - 1 and number == 1))
            c.border = BOX
            c.alignment = Alignment(horizontal=col[3], vertical="center", wrap_text=True)
            if col[2]:
                # В строках ноль печатается (отход 0,0 кг); скрывается только в пустом бланке — формулой "" и в итогах.
                c.number_format = col[2] if col[4] else col[2].split(";")[0]
            if isinstance(value, str) and not value.startswith("="):
                height = max(height, text_height(value, col[1], 9, factor=0.9))
        ws.row_dimensions[r].height = height
    end = first + len(rows) - 1

    row = end + 1
    total = spec["total"]
    if total:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total["span"])
        for i in range(1, n + 1):
            c = ws.cell(row=row, column=i)
            c.fill = GREY_FILL
            c.border = BOX
            c.font = Font(name=FONT, bold=True, size=9)
            c.alignment = Alignment(horizontal="right" if i > total["span"] else "left", vertical="center", wrap_text=True)
            if cols[i - 1][2]:
                c.number_format = cols[i - 1][2]
        ws.cell(row=row, column=1, value=total["label"])
        for i, formula in total["cells"].items():
            letter = get_column_letter(i)
            ws.cell(row=row, column=i, value=formula.format(c=f"{letter}{first}:{letter}{end}", a=first, b=end))
        for i, value in (section_data.get("total") or {}).items():
            ws.cell(row=row, column=int(i), value=value)
        spans = [(total["label"], sum(c[1] for c in cols[:total["span"]]))]
        spans += [(v, cols[int(i) - 1][1]) for i, v in (section_data.get("total") or {}).items() if isinstance(v, str)]
        ws.row_dimensions[row].height = max([18] + [text_height(v, w - 1, 9) for v, w in spans if not v.startswith("=")])
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
    ws.cell(row=row, column=1, value="Технологические указания и требования участка:").font = Font(name=FONT, bold=True, size=9)
    row += 1
    for note in section_data.get("notes") or spec["notes"]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n)
        c = ws.cell(row=row, column=1, value="•  " + note)
        c.font = Font(name=FONT, size=9, color="303030")
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.row_dimensions[row].height = text_height("•  " + note, total_width * 1.2, 9)
        row += 1

    row += 1
    dash = "_" * (20 if len(spec["signers"]) == 1 else 12)
    signers = [("Нач. КТО:", '="' + dash + ' / "&Данные!B4')] + \
              [(title + ":", dash + " / " + dash) for title in spec["signers"]]
    bounds = [round(n * k / len(signers)) for k in range(len(signers) + 1)]
    for k, (title_text, line) in enumerate(signers):
        a, b = bounds[k] + 1, bounds[k + 1]
        for r, value, font in ((row, title_text, Font(name=FONT, bold=True, size=9)), (row + 1, line, Font(name=FONT, size=9))):
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            c = ws.cell(row=r, column=a, value=value)
            c.font = font
            c.alignment = Alignment(horizontal="left", vertical="bottom")
        ws.row_dimensions[row + 1].height = 22
    last_row = row + 1

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.6
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "1:7"
    ws.print_area = f"A1:{last}{last_row}"
    ws.oddFooter.right.text = "&A, стр. &P из &N"
    ws.oddFooter.right.size = 8
    ws.protection.set_password(PROD_PASSWORD)
    ws.protection.sheet = True
    return ws


def prod_note(data=None):
    """Пустой бланк (data=None) или заполненная заявка из словаря той же структуры, что читает кнопка К-4."""
    data = data or {}
    wb = Workbook()
    ds = wb.active
    ds.title = "Данные"
    for r, (label, default) in enumerate(DATA_FIELDS, start=1):
        ds.cell(row=r, column=1, value=label).font = Font(name=FONT, bold=True, size=10)
        ds.cell(row=r, column=2, value=data.get(label, default)).font = Font(name=FONT, size=10)
    ds.column_dimensions["A"].width = 14
    ds.column_dimensions["B"].width = 40
    ds.protection.set_password(PROD_PASSWORD)
    ds.protection.sheet = True
    sections = data.get("Разделы") or {}
    for spec in PROD_SECTIONS:
        prod_sheet(wb, spec, data, sections.get(spec["sheet"], {}))
    ds.sheet_state = "hidden"
    wb.active = 1
    wb.security = WorkbookProtection(workbookPassword=PROD_PASSWORD, lockStructure=True)
    wb.calculation.fullCalcOnLoad = True
    return wb


def change_log():
    wb = Workbook()
    ws = wb.active
    ws.title = "Изменения"
    title(ws, "Журнал изменений изделия", [("Изделие (шифр):", ""), ("Папка:", "")])
    columns = [("№", 6), ("Ревизия", 9), ("Дата", 12), ("Кто", 20), ("Документ", 34), ("Что изменено", 44), ("Причина", 34),
               ("Код причины", 12), ("Задел", 14), ("Применяемость", 30)]
    header_row(ws, 5, columns)
    body_rows(ws, 6, 30, len(columns))
    dv = DataValidation(type="list", formula1='"использовать,доработать,в брак"', allow_blank=True)
    dv.add("I6:I500")
    ws.add_data_validation(dv)
    ws["A37"] = "Строки дописывают кнопки «Новая ревизия» и «Снимок эталона»; вручную журнал не ведётся."
    ws["A37"].font = Font(name=FONT, italic=True, size=9)
    return wb


def main():
    if len(sys.argv) == 4 and sys.argv[1] == "--prod":
        data = json.loads(Path(sys.argv[2]).read_text(encoding="utf-8"))
        prod_note(data).save(sys.argv[3])
        print("записан", sys.argv[3])
        return 0
    OUT.mkdir(parents=True, exist_ok=True)
    for name, build in (("СЗ_заказа_макет.xlsx", order_note), ("СЗ_на_производство.xlsx", prod_note), ("Изменения.xlsx", change_log)):
        path = OUT / name
        build().save(path)
        print("записан", path)
    NORMS_PATH.parent.mkdir(parents=True, exist_ok=True)
    norms_workbook().save(NORMS_PATH)
    print("записан", NORMS_PATH)
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
